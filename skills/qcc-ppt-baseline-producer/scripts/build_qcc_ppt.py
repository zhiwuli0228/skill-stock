#!/usr/bin/env python3
"""Build a QCC PPTX from qcc-data.yaml using the supplied template page-type layouts.

V2.4 rule: the template owns background, brand layer, footer/logo and page
visual identity. This generator must not create a new visual system. It maps
the approved template pages to page roles, sanitizes inherited template placeholders/static text,
prevents duplicate template-owned end content, and only fills the safe content area.
"""
from __future__ import annotations
import argparse, yaml
from pathlib import Path
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_VERTICAL_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR

W_IN, H_IN = 13.333333, 7.5
FOOTER_MODE = "none"  # none | minimal | full. Default: template owns footer/logo.
TEMPLATE_LAYOUTS = {}  # populated by configure_template_layouts(prs)
C = {
    "black": "231815", "red": "C8102E", "orange": "EB5C01", "green": "009944", "blue": "1F4E79",
    "purple": "7030A0", "grey": "898989", "line": "DDDDDD", "light": "F7F7F7", "white": "FFFFFF", "red_light": "FDECEE", "orange_light": "FFF2E8", "blue_light": "EAF3FF", "green_light": "EAF7EA", "purple_light": "F3ECFA"
}


def rgb(hexstr: str) -> RGBColor:
    h = hexstr.strip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def load_yaml(path: Path) -> dict:
    with path.open("r", encoding="utf-8") as f:
        return yaml.safe_load(f) or {}


def add_text(slide, text, x, y, w, h, size=12, color="black", bold=False, align="left"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}.get(align, PP_ALIGN.LEFT)
    run = p.add_run()
    run.text = str(text)
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(C.get(color, color))
    return box


def shape(slide, x, y, w, h, fill="white", line="line", radius=True):
    st = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(st, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = rgb(C.get(fill, fill))
    shp.line.color.rgb = rgb(C.get(line, line)); shp.line.width = Pt(0.7)
    return shp


def line(slide, x1, y1, x2, y2, color="grey", width=1.0):
    shp = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shp.line.color.rgb = rgb(C.get(color, color)); shp.line.width = Pt(width)
    return shp


def arrow(slide, x1, y1, x2, y2, color="grey"):
    shp = line(slide, x1, y1, x2, y2, color, 1.1)
    try:
        shp.line.end_arrowhead = True
    except Exception:
        pass
    return shp


def header(slide, title, subtitle=None):
    add_text(slide, title, 0.95, 0.62, 11.6, 0.43, 25, "black", True)
    line(slide, 0.98, 1.16, 1.93, 1.16, "red", 2.0)
    if subtitle:
        add_text(slide, subtitle, 0.98, 1.28, 11.0, 0.20, 9, "grey")


def footer(slide, page):
    """Optional generated footer.

    Default is ``none`` because approved corporate templates normally already
    contain page-branding and footer objects in their slide layouts. Adding
    generated footer/logo elements on top of those layouts is forbidden in
    template-driven mode.
    """
    if FOOTER_MODE == "none":
        return
    add_text(slide, str(page), 0.95, 7.13, 0.25, 0.13, 8, "black")
    if FOOTER_MODE == "minimal":
        add_text(slide, "Huawei Confidential", 1.35, 7.13, 1.8, 0.13, 8, "black")
    elif FOOTER_MODE == "full":
        raise RuntimeError("footer-mode=full is disabled in v2.4 template-driven mode; use an approved blank template profile instead")




def _delete_shape(shp):
    """Remove a shape from a slide XML tree."""
    el = shp._element
    el.getparent().remove(el)


PROMPT_FRAGMENTS = [
    "单击此处添加标题", "单击此处添加文本", "Click to add title",
    "Click to add text", "Click to add subtitle", "添加标题", "添加文本",
]


def _shape_text(shp) -> str:
    if not getattr(shp, "has_text_frame", False):
        return ""
    try:
        return "\n".join(p.text for p in shp.text_frame.paragraphs).strip()
    except Exception:
        return ""


def _contains_prompt_text(shp) -> bool:
    txt = _shape_text(shp)
    return bool(txt and any(frag in txt for frag in PROMPT_FRAGMENTS))


def clear_template_placeholders(slide):
    """Remove slide-level editable placeholders created by the template layout.

    V2.4 also sanitizes layout-level placeholders before slide generation. This
    function remains as a second guard for placeholders materialized on the
    concrete slide.
    """
    to_remove = []
    for shp in list(slide.shapes):
        try:
            if getattr(shp, "is_placeholder", False) or _contains_prompt_text(shp):
                to_remove.append(shp)
        except Exception:
            continue
    for shp in to_remove:
        _delete_shape(shp)


def sanitize_template_layouts():
    """Remove inherited editable placeholders/static template text that would
    overlap generated content.

    The important distinction is:
    - cover/agenda/body/chart layouts are used as clean branded backgrounds;
      all editable placeholders and duplicate static title text are removed.
    - end layout is template-owned; it keeps its static Thank-you/legal content,
      and the generator must not add another Thank-you block.
    """
    seen_layouts = set()
    seen_masters = set()
    for role, layout in TEMPLATE_LAYOUTS.items():
        # First sanitize master placeholders/prompts. Keep non-placeholder static
        # branding/legal text because the template owns those objects.
        master = getattr(layout, "slide_master", None)
        if master is not None and id(master) not in seen_masters:
            seen_masters.add(id(master))
            master_remove = []
            for shp in list(master.shapes):
                is_ph = bool(getattr(shp, "is_placeholder", False))
                if is_ph or _contains_prompt_text(shp):
                    master_remove.append(shp)
            for shp in master_remove:
                _delete_shape(shp)

        if id(layout) in seen_layouts:
            continue
        seen_layouts.add(id(layout))
        to_remove = []
        for shp in list(layout.shapes):
            txt = _shape_text(shp)
            is_ph = bool(getattr(shp, "is_placeholder", False))
            # Editable placeholders are never allowed to remain on generated pages.
            if is_ph or _contains_prompt_text(shp):
                to_remove.append(shp)
                continue
            # Agenda layout contains a static "目录" text object. Since agenda()
            # writes a generated agenda title, remove the inherited duplicate.
            if role == "agenda" and txt.strip() in {"目录", "Contents", "CONTENTS"}:
                to_remove.append(shp)
                continue
            # Non-end layouts must not own visible text; the generator owns slide
            # titles/body. Keep pictures, lines and other decorative objects.
            if role != "end" and txt.strip():
                to_remove.append(shp)
                continue
        for shp in to_remove:
            _delete_shape(shp)

def remove_template_seed_slides(prs):
    """Remove sample/template seed slides after their layouts have been mapped.

    This does not delete layouts or masters. It only removes the seed pages that
    ship with the template, so the generated deck starts at page 1 while still
    using the template-owned page types.
    """
    while len(prs.slides) > 0:
        rId = prs.slides._sldIdLst[0].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[0]


def configure_template_layouts(prs):
    """Map approved template layouts to semantic page roles.

    Prefer the bundled seed slides because the template uses multiple masters;
    `prs.slide_layouts` only exposes layouts from the first master in
    python-pptx. Falling back to `prs.slide_layouts[min(...)]` is exactly what
    caused content slides to use the wrong/end background.
    """
    seed_slides = list(prs.slides)
    if len(seed_slides) >= 5:
        TEMPLATE_LAYOUTS.clear()
        TEMPLATE_LAYOUTS.update({
            "cover": seed_slides[0].slide_layout,
            "agenda": seed_slides[1].slide_layout,
            "body": seed_slides[2].slide_layout,
            "chart": seed_slides[3].slide_layout,
            "end": seed_slides[4].slide_layout,
        })
        return

    # Fallback for a custom template with no seed pages: scan all masters, not
    # only prs.slide_layouts from the first master.
    layouts = []
    for master in prs.slide_masters:
        layouts.extend(list(master.slide_layouts))
    if not layouts:
        raise RuntimeError("template has no slide layouts; cannot build template-driven deck")

    def by_name(keys, fallback):
        keys_l = [k.lower() for k in keys]
        for layout in layouts:
            name = (layout.name or "").lower()
            if any(k in name for k in keys_l):
                return layout
        return layouts[min(fallback, len(layouts)-1)]

    TEMPLATE_LAYOUTS.clear()
    TEMPLATE_LAYOUTS.update({
        "cover": by_name(["探索", "cover", "title"], 0),
        "agenda": by_name(["contents", "目录", "agenda"], 1),
        "body": by_name(["chinese", "正文", "text", "body"], 2),
        "chart": by_name(["chart", "图表", "chinese"], 2),
        "end": by_name(["end", "结束"], 3),
    })


def slide(prs, page=None, title=None, subtitle=None, role=None):
    if role is None:
        if page == 2 and title is None:
            role = "agenda"
        elif title is None and page is None:
            role = "cover"
        else:
            role = "body"
    layout = TEMPLATE_LAYOUTS.get(role) or TEMPLATE_LAYOUTS.get("body") or prs.slide_layouts[0]
    s = prs.slides.add_slide(layout)
    clear_template_placeholders(s)
    if title:
        header(s, title, subtitle)
    if page is not None:
        footer(s, page)
    return s


def card(slide, x, y, w, h, title, body, color="red"):
    """Draw a content card.

    v2.5 safety rule: never place body text below the card when the card is
    compact. The old implementation always used ``y+0.52`` for body text and
    ``h-0.62`` for body height; any card below 0.62in generated a negative text
    box and caused visible overlap/escape after PPT/PDF rendering.
    """
    shape(slide, x, y, w, h, "white", "line")
    shape(slide, x, y, 0.06, h, color, color, False)
    if h < 0.72:
        # Compact row: title and body share one baseline inside the row.
        title_w = min(3.25, max(1.25, w * 0.30))
        add_text(slide, title, x+0.22, y+h/2-0.09, title_w, 0.18, 10.0, color, True)
        add_text(slide, body, x+0.32+title_w, y+0.10, w-title_w-0.55, h-0.20, 8.4, "black")
    else:
        add_text(slide, title, x+0.18, y+0.15, w-0.35, 0.20, 12, color, True)
        add_text(slide, body, x+0.18, y+0.48, w-0.35, max(0.18, h-0.56), 9.2, "black")


def kpi(slide, x, y, w, h, num, label, desc="", color="red"):
    shape(slide, x, y, w, h, "white", "line")
    add_text(slide, num, x+0.08, y+0.18, w-0.16, 0.30, 22, color, True, "center")
    add_text(slide, label, x+0.08, y+0.55, w-0.16, 0.16, 9.5, "black", True, "center")
    if desc:
        add_text(slide, desc, x+0.08, y+0.78, w-0.16, 0.15, 8, "grey", False, "center")


def conclusion(slide, text):
    if not text: return
    shape(slide, 0.95, 6.32, 11.45, 0.26, "red_light", "red")
    add_text(slide, text, 1.08, 6.39, 11.15, 0.08, 8.6, "red", True, "center")



# -----------------------------
# v2.7 rich visual primitives
# -----------------------------

def pill(slide, x, y, w, h, text, color="red", fill="white", size=9.2, bold=True):
    shape(slide, x, y, w, h, fill, color)
    add_text(slide, text, x+0.08, y+0.08, w-0.16, max(0.12, h-0.16), size, color, bold, "center")


def rich_table(slide, x, y, widths, row_h, headers, rows, accent_cols=None, max_rows=None):
    """Template-safe compact table with a light header and alternating body rows."""
    if max_rows is not None:
        rows = rows[:max_rows]
    total_w=sum(widths)
    shape(slide, x, y, total_w, row_h, "red_light", "red", False)
    cx=x
    for i,h in enumerate(headers):
        add_text(slide, h, cx+0.03, y+0.10, widths[i]-0.06, row_h-0.18, 8.6, "red", True, "center")
        if i>0:
            line(slide, cx, y+0.06, cx, y+row_h-0.06, "white", 0.8)
        cx += widths[i]
    for r,row in enumerate(rows):
        yy=y+row_h+(r*row_h)
        fill="white" if r%2==0 else "light"
        shape(slide, x, yy, total_w, row_h, fill, "line", False)
        cx=x
        for c,v in enumerate(row):
            col="red" if accent_cols and c in accent_cols else "black"
            bold=bool(accent_cols and c in accent_cols)
            add_text(slide, v, cx+0.04, yy+0.10, widths[c]-0.08, row_h-0.18, 8.2, col, bold, "center" if c>0 else "left")
            if c>0:
                line(slide, cx, yy+0.06, cx, yy+row_h-0.06, "white", 0.6)
            cx += widths[c]



def _candidate_name(c: dict) -> str:
    return str(c.get("name") or c.get("topic") or c.get("candidate") or c.get("title") or "待补充")


def _criterion(c: dict, key: str, alt: str | None = None) -> int:
    if key in c and c.get(key) not in (None, ""):
        try: return int(c.get(key))
        except Exception: return 0
    criteria = c.get("criteria", {}) or {}
    if key in criteria and criteria.get(key) not in (None, ""):
        try: return int(criteria.get(key))
        except Exception: return 0
    if alt and alt in criteria and criteria.get(alt) not in (None, ""):
        try: return int(criteria.get(alt))
        except Exception: return 0
    return 0


def _candidate_score(c: dict) -> int:
    if c.get("score") not in (None, ""):
        try: return int(c.get("score"))
        except Exception: pass
    return sum(_criterion(c, k, alt) for k, alt in [
        ("importance", None), ("urgency", None), ("feasibility", None),
        ("data_availability", None), ("promotion_value", "promotability")
    ])


def normalize_candidates(candidates: list[dict]) -> list[dict]:
    out=[]
    for c in candidates or []:
        row=dict(c)
        row["_name"]=_candidate_name(c)
        row["_score"]=_candidate_score(c)
        row["importance"]=_criterion(c,"importance")
        row["urgency"]=_criterion(c,"urgency")
        row["feasibility"]=_criterion(c,"feasibility")
        row["data_availability"]=_criterion(c,"data_availability")
        row["promotion_value"]=_criterion(c,"promotion_value","promotability")
        out.append(row)
    return sorted(out, key=lambda x: x.get("_score",0), reverse=True)


def score_matrix(slide, x, y, candidates, with_top_summary=False):
    """QCC checklist / scoring matrix with safe TOP-summary formatting.

    Supports both v2.8 flat fields and example.yaml's nested `criteria` format.
    """
    candidates=normalize_candidates(candidates)[:4]
    headers=["候选课题","重要性","紧迫性","可行性","数据","推广","总分"]
    widths=[3.55,0.82,0.82,0.82,0.82,0.82,0.82]
    rows=[]
    for c in candidates:
        rows.append([c.get("_name",""),c.get("importance",""),c.get("urgency",""),c.get("feasibility",""),c.get("data_availability",""),c.get("promotion_value",""),c.get("_score","")])
    rich_table(slide,x,y,widths,0.46,headers,rows,accent_cols={6},max_rows=4)
    if with_top_summary and candidates:
        ranking_summary(slide, 9.40, y, 2.55, 1.96, candidates[:3], title="TOP 课题摘要")


def ranking_summary(slide, x, y, w, h, candidates, title="TOP 摘要"):
    """Formal-review-safe TOP ranking side card.

    Fixed zones prevent two-digit scores such as 24/20/18 from vertical wrapping.
    """
    shape(slide, x, y, w, h, "white", "line")
    add_text(slide, title, x+0.18, y+0.18, w-0.36, 0.16, 9.4, "red", True)
    row_y=y+0.50
    colors=["red","orange","blue"]
    for i,c in enumerate((candidates or [])[:3]):
        yy=row_y+i*0.46
        shape(slide, x+0.18, yy, w-0.36, 0.38, "white", "line", False)
        shape(slide, x+0.18, yy, 0.36, 0.38, colors[i%3], colors[i%3], False)
        add_text(slide, str(i+1), x+0.29, yy+0.13, 0.12, 0.08, 8.2, "white", True, "center")
        add_text(slide, c.get("_name", _candidate_name(c)), x+0.64, yy+0.07, w-1.48, 0.20, 8.0, "black")
        add_text(slide, f"{c.get('_score', _candidate_score(c))}分", x+w-0.72, yy+0.12, 0.48, 0.09, 8.6, colors[i%3], True, "right")

def flow_bar(slide, x, y, steps, box_w=1.78, box_h=1.30, gap=0.42):
    """Horizontal flow with numbered badges and arrow connectors."""
    for i,(title, body, color) in enumerate(steps):
        xx=x+i*(box_w+gap)
        shape(slide, xx, y, box_w, box_h, "white", "line")
        shape(slide, xx, y, box_w, 0.30, "red_light", color, False)
        add_text(slide, f"{i+1:02d}", xx+0.08, y+0.08, 0.35, 0.10, 8.6, color, True, "center")
        add_text(slide, title, xx+0.45, y+0.07, box_w-0.55, 0.12, 8.8, color, True)
        add_text(slide, body, xx+0.14, y+0.48, box_w-0.28, box_h-0.58, 8.2, "black")
        if i < len(steps)-1:
            arrow(slide, xx+box_w+0.05, y+box_h/2, xx+box_w+gap-0.10, y+box_h/2, "grey")


def vertical_timeline(slide, x, y, items, step_h=0.70):
    """Vertical milestone/timeline rows for roadmap pages."""
    line(slide, x+0.22, y+0.20, x+0.22, y+step_h*len(items)-0.10, "red", 1.5)
    for i,(title, body, color) in enumerate(items):
        yy=y+i*step_h
        shape(slide, x+0.04, yy+0.09, 0.36, 0.36, color, color)
        add_text(slide, str(i+1), x+0.13, yy+0.20, 0.18, 0.08, 8.4, "white", True, "center")
        shape(slide, x+0.58, yy, 10.95, step_h-0.10, "white", "line")
        add_text(slide, title, x+0.78, yy+0.14, 2.25, 0.13, 9.4, color, True)
        add_text(slide, body, x+3.0, yy+0.14, 8.15, 0.13, 8.4, "black")


def fishbone(slide, x, y, items, conclusion_text):
    """Light fishbone-style cause map. Keeps text inside safe area."""
    # main spine
    line(slide, x+0.35, y+2.0, x+10.0, y+2.0, "grey", 1.5)
    arrow(slide, x+10.0, y+2.0, x+10.8, y+2.0, "red")
    shape(slide, x+10.85, y+1.62, 1.15, 0.76, "red_light", "red")
    add_text(slide, "关键\n治理缺口", x+11.03, y+1.80, 0.80, 0.18, 8.5, "red", True, "center")
    colors=["red","orange","green","blue","purple","red"]
    anchors=[(1.0,0.52,2.05,1.25),(4.1,0.52,5.0,1.25),(7.2,0.52,7.95,1.25),(1.0,3.00,2.05,2.75),(4.1,3.00,5.0,2.75),(7.2,3.00,7.95,2.75)]
    for i,it in enumerate(items[:6]):
        bx,by,lx,ly=anchors[i]
        c=colors[i%len(colors)]
        # branch
        line(slide, x+lx, y+ly, x+lx+0.9, y+2.0, c, 1.1)
        shape(slide, x+bx, y+by, 2.45, 0.78, "white", "line")
        shape(slide, x+bx, y+by, 0.06, 0.78, c, c, False)
        add_text(slide, it.get("dimension",""), x+bx+0.18, y+by+0.12, 2.05, 0.11, 8.8, c, True)
        add_text(slide, it.get("cause",""), x+bx+0.18, y+by+0.38, 2.05, 0.18, 8.0, "black")
    if conclusion_text:
        card(slide, x+0.15, y+4.15, 11.55, 0.65, "收敛判断", conclusion_text, "red")


def mapping_matrix(slide, x, y, items):
    headers=["关键要因","对策","实施方式","验证方式"]
    widths=[2.35,3.2,3.0,2.8]
    rows=[]
    for c in items[:5]:
        rows.append([c.get("key_factor",""), c.get("countermeasure",""), c.get("implementation",""), c.get("validation","")])
    rich_table(slide,x,y,widths,0.55,headers,rows,accent_cols={0},max_rows=5)


def gate_grid(slide, x, y, guards):
    colors=["red","orange","blue","green","red","orange","blue","green"]
    for i,g in enumerate(guards[:8]):
        xx=x+(i%4)*2.82; yy=y+(i//4)*1.24
        shape(slide, xx, yy, 2.42, 0.92, "white", "line")
        shape(slide, xx, yy, 0.42, 0.92, colors[i], colors[i], False)
        add_text(slide, g[0], xx+0.08, yy+0.34, 0.26, 0.12, 9, "white", True, "center")
        add_text(slide, g[1], xx+0.56, yy+0.18, 1.62, 0.12, 9.0, colors[i], True)
        add_text(slide, g[2], xx+0.56, yy+0.50, 1.66, 0.12, 8.0, "black")


def pdca_cycle(slide, x, y):
    items=[("P","问题与目标","red"),("D","对策实施","orange"),("C","效果确认","green"),("A","标准推广","blue")]
    coords=[(x+0.9,y),(x+4.0,y),(x+4.0,y+1.85),(x+0.9,y+1.85)]
    for i,(letter,title,color) in enumerate(items):
        xx,yy=coords[i]
        shape(slide, xx, yy, 2.15, 1.05, "white", "line")
        shape(slide, xx+0.18, yy+0.18, 0.55, 0.55, color, color)
        add_text(slide, letter, xx+0.31, yy+0.34, 0.30, 0.10, 12, "white", True, "center")
        add_text(slide, title, xx+0.88, yy+0.38, 1.05, 0.12, 9.2, color, True)
    arrow(slide,x+3.08,y+0.52,x+3.85,y+0.52,"grey")
    arrow(slide,x+5.08,y+1.08,x+5.08,y+1.80,"grey")
    arrow(slide,x+3.85,y+2.38,x+3.08,y+2.38,"grey")
    arrow(slide,x+0.93,y+1.80,x+0.93,y+1.08,"grey")


# -----------------------------
# v2.9 QCC mandatory method-page primitives
# -----------------------------

def _topic_selection(data: dict) -> dict:
    return data.get("qcc_process", {}).get("topic_selection", {}) or {}


def _current_state(data: dict) -> dict:
    return data.get("qcc_process", {}).get("current_state", {}) or {}


def _method_data(data: dict, key: str, default=None):
    return (data.get("qcc_methods", {}) or {}).get(key, default)


def _short(text, n=38):
    s=str(text or "").replace("\n", " ").strip()
    return s if len(s)<=n else s[:n-1]+"…"


def brainstorm_page(prs, data, page):
    topic=_topic_selection(data)
    candidates=normalize_candidates(topic.get("candidates", []))[:6]
    if not candidates:
        candidates=[{"_name": x, "description":"待补充"} for x in ["候选方向1","候选方向2","候选方向3","候选方向4","候选方向5","候选方向6"]]
    s=slide(prs,page,"主题评审｜头脑风暴", "围绕课题痛点收集候选改善方向；本页只发散，不评分。")
    # Formal-review-safe pattern: left anchor + strict 2×3 idea grid. No connector clutter.
    shape(s,0.98,1.85,2.05,2.70,"red_light","red")
    add_text(s,"发散主题",1.18,2.08,1.62,0.16,11.0,"red",True,"center")
    center=topic.get("problem_statement") or topic.get("reason") or data.get("circle",{}).get("topic") or data.get("meta",{}).get("report_title","QCC课题")
    add_text(s,_short(center,52),1.20,2.48,1.60,0.72,10.0,"black",True,"center")
    for i,chip in enumerate(["不评分","不排序","先发散"]):
        pill(s,1.16+i*0.58,4.18,0.50,0.25,chip,"red","white",8.0,True)
    colors=["red","orange","blue","green","purple","red"]
    for i,c in enumerate(candidates[:6]):
        xx=3.55+(i%3)*2.72; yy=1.72+(i//3)*1.46
        shape(s,xx,yy,2.36,1.04,"white","line")
        shape(s,xx,yy,0.38,1.04,colors[i],colors[i],False)
        add_text(s,str(i+1),xx+0.13,yy+0.43,0.12,0.09,9.2,"white",True,"center")
        add_text(s,_short(c.get("_name",_candidate_name(c)),22),xx+0.55,yy+0.20,1.58,0.18,9.4,"black",True)
        desc=c.get("description") or c.get("reason") or c.get("pain_point") or "候选改善方向"
        add_text(s,_short(desc,24),xx+0.55,yy+0.58,1.58,0.14,8.0,"grey")
    card(s,1.0,5.35,11.30,0.62,"发散输出","候选方向进入亲和图聚类，再通过检查表/矩阵评分选择改善课题。","red")


def affinity_page(prs, data, page):
    topic=_topic_selection(data)
    candidates=normalize_candidates(topic.get("candidates", []))[:6]
    groups=_method_data(data,"affinity_groups")
    if not groups:
        names=[c.get("_name",_candidate_name(c)) for c in candidates] or ["候选方向1","候选方向2","候选方向3","候选方向4","候选方向5","候选方向6"]
        groups=[
            {"label":"稳定性治理","ideas":names[0:2] or ["待补充"]},
            {"label":"效率优化","ideas":names[2:4] or ["待补充"]},
            {"label":"过程治理","ideas":names[4:6] or ["待补充"]},
        ]
    s=slide(prs,page,"主题评审｜亲和图", "将头脑风暴候选项按问题属性聚类，形成可比较的课题方向。")
    colors=[("red","red_light"),("orange","orange_light"),("blue","blue_light")]
    for i,g in enumerate(groups[:3]):
        x=1.0+i*3.78; y=1.72
        color,fill=colors[i%3]
        shape(s,x,y,3.25,3.48,fill,color)
        add_text(s,g.get("label",f"聚类{i+1}"),x+0.20,y+0.20,2.80,0.18,11,color,True,"center")
        for j,idea in enumerate((g.get("ideas") or [])[:3]):
            yy=y+0.68+j*0.78
            shape(s,x+0.28,yy,2.70,0.54,"white","line")
            add_text(s,_short(idea,28),x+0.46,yy+0.18,2.34,0.10,8.5,"black",False,"center")
    card(s,1.0,5.55,11.3,0.55,"聚类结论","候选方向集中在稳定性治理、效率优化和过程治理三类，进入检查表评分选择。","red")


def checklist_topic_page(prs, data, page):
    topic=_topic_selection(data)
    s=slide(prs,page,"主题评审｜检查表：矩阵评分选择改善课题", "静态线程池无法同时兼顾稳定性、资源效率和响应速度。")
    cand=topic.get("candidates", [])[:4]
    score_matrix(s,1.00,1.72,cand,with_top_summary=True)
    card(s,1.0,5.52,11.30,0.55,"选择结论","优先选择高重要性、高紧迫性、具备数据验证条件且具备推广价值的课题。","red")


def sipoc_page(prs, data, page):
    current=_current_state(data)
    sipoc=_method_data(data,"sipoc") or current.get("sipoc") or {}
    headers=["Supplier","Input","Process","Output","Customer"]
    if not sipoc:
        sipoc={"supplier":"业务流量/运维配置","input":"线程池参数/运行指标","process":["采样","诊断","决策","执行","审计"],"output":"治理建议/调整结果","customer":"服务运行与运维人员"}
    row=[sipoc.get("supplier","待补充"), sipoc.get("input","待补充"), "\n".join(sipoc.get("process",[]) if isinstance(sipoc.get("process"),list) else [str(sipoc.get("process","待补充"))]), sipoc.get("output","待补充"), sipoc.get("customer","待补充")]
    s=slide(prs,page,"把握现状｜SIPOC", "用 SIPOC 明确问题链路边界，避免现状调查失焦。")
    rich_table(s,0.95,1.86,[2.10,2.10,2.90,2.10,2.10],0.74,headers,[row],accent_cols={2},max_rows=1)
    # Add process mini-flow for readability.
    proc=sipoc.get("process",[]) if isinstance(sipoc.get("process"),list) else ["采样","诊断","决策","执行","审计"]
    flow_bar(s,1.35,3.62,[(p,"",["red","orange","blue","green","purple"][i%5]) for i,p in enumerate(proc[:5])],box_w=1.65,box_h=0.70,gap=0.35)
    card(s,1.0,5.52,11.30,0.55,"现状边界","后续柏拉图与子流程图必须围绕 SIPOC 识别的问题链路展开。","red")


def pareto_page(prs, data, page):
    current=_current_state(data)
    items=_method_data(data,"pareto_items") or current.get("pareto_items")
    if not items:
        problems=current.get("problems",[])[:5]
        items=[{"name":p.get("scenario",f"问题{i+1}"),"value":max(1,5-i)} for i,p in enumerate(problems)] or [{"name":"待补充问题A","value":5},{"name":"待补充问题B","value":3},{"name":"待补充问题C","value":2}]
    items=sorted(items,key=lambda x: float(x.get("value",x.get("count",0)) or 0),reverse=True)[:5]
    vals=[float(it.get("value",it.get("count",0)) or 0) for it in items]
    total=sum(vals) or 1
    s=slide(prs,page,"把握现状｜柏拉图：识别关键 80% 改进项", "按影响度/发生次数降序识别少数关键问题。")
    chart_x,chart_y,chart_w,chart_h=1.10,1.88,7.85,3.35
    line(s,chart_x,chart_y+chart_h,chart_x+chart_w,chart_y+chart_h,"grey",1.0)
    line(s,chart_x,chart_y,chart_x,chart_y+chart_h,"grey",1.0)
    maxv=max(vals) or 1
    cum=0
    prev=None
    bar_w=min(0.82, chart_w/(len(items)*1.4))
    for i,(it,v) in enumerate(zip(items,vals)):
        xx=chart_x+0.55+i*(chart_w-1.0)/max(1,len(items)-1)
        bh=chart_h*(v/maxv)*0.86
        shape(s,xx,chart_y+chart_h-bh,bar_w,bh,"red","red",False)
        add_text(s,str(int(v)),xx-0.02,chart_y+chart_h-bh-0.25,bar_w+0.04,0.10,8.0,"red",True,"center")
        add_text(s,_short(it.get("name","问题"),8),xx-0.12,chart_y+chart_h+0.10,bar_w+0.25,0.22,8.0,"black",False,"center")
        cum += v
        cy=chart_y+chart_h-(cum/total)*chart_h*0.86
        cx=xx+bar_w/2
        shape(s,cx-0.035,cy-0.035,0.07,0.07,"blue","blue",False)
        if prev:
            line(s,prev[0],prev[1],cx,cy,"blue",1.1)
        prev=(cx,cy)
        add_text(s,f"{cum/total:.0%}",cx-0.20,cy-0.28,0.40,0.08,8.0,"blue",False,"center")
    # 80% reference.
    y80=chart_y+chart_h-chart_h*0.86*0.80
    line(s,chart_x,y80,chart_x+chart_w,y80,"orange",1.0)
    add_text(s,"80%",chart_x+chart_w+0.10,y80-0.06,0.40,0.08,8.0,"orange",True)
    top_names="、".join(_short(i.get("name",""),10) for i in items[:2])
    card(s,9.45,1.90,2.55,2.15,"关键 80%",f"优先关注：\n{top_names}\n\n其余问题进入后续跟踪。","red")
    card(s,1.0,5.52,11.30,0.55,"现状判断","少数关键问题贡献主要影响，后续子流程图聚焦关键异常链路。","red")


def subprocess_page(prs, data, page):
    current=_current_state(data)
    steps=_method_data(data,"subprocess_steps") or current.get("subprocess_steps") or ["配置输入","运行采样","问题识别","人工判断","调整执行","结果复盘"]
    pain=_method_data(data,"subprocess_pain_points") or current.get("subprocess_pain_points") or ["发现滞后","经验依赖","缺少审计"]
    s=slide(prs,page,"把握现状｜子流程图", "把关键问题放回实际流程，定位异常点和改善入口。")
    step_items=[]
    for i,st in enumerate(steps[:6]):
        note="异常点" if i in {2,3,4} else ""
        step_items.append((st,note,["red","orange","blue","green","purple","red"][i%6]))
    flow_bar(s,0.95,2.10,step_items,box_w=1.55,box_h=1.18,gap=0.34)
    for i,p in enumerate(pain[:3]):
        pill(s,1.20+i*3.35,4.10,2.35,0.36,f"痛点{i+1}：{_short(p,14)}",["red","orange","blue"][i],"white",8.0,True)
    card(s,1.0,5.42,11.30,0.65,"流程结论","异常集中在问题识别、判断决策和调整复盘环节，需要进入根因分析。","red")


def cause_matrix_page(prs, data, page):
    q=data.get("qcc_process",{})
    factors=q.get("key_factor_confirmation",{}).get("factors",[])[:5]
    if not factors:
        dims=q.get("cause_analysis",{}).get("dimensions",[])[:5]
        factors=[{"cause":d.get("cause") or d.get("detail") or d.get("name") or d.get("dimension","待补充"),"impact":4,"frequency":4,"control":4,"evidence":4} for d in dims]
    rows=[]
    for f in factors[:5]:
        vals=[]
        for k in ["impact","frequency","controllability","evidence_availability"]:
            try: vals.append(int(f.get(k,4)))
            except Exception: vals.append(4)
        score=sum(vals)
        keep="保留" if score>=14 else "观察"
        rows.append([f.get("cause",""), vals[0], vals[1], vals[2], vals[3], score, keep])
    s=slide(prs,page,"根因分析｜矩阵图：筛选关键要因", "鱼骨图发散后，使用矩阵图按影响、频次、可控性和证据条件筛选。")
    rich_table(s,0.95,1.78,[3.00,0.78,0.78,0.88,0.88,0.72,0.82],0.48,["疑似原因","影响","频次","可控","证据","分数","判断"],rows,accent_cols={0,5},max_rows=5)
    card(s,1.0,5.52,11.30,0.55,"矩阵结论","优先保留高影响、高频次、可控且具备验证证据的疑似原因，进入根因验证。","red")


def root_validation_page(prs, data, page):
    q=data.get("qcc_process",{})
    factors=q.get("key_factor_confirmation",{}).get("factors",[])[:4]
    s=slide(prs,page,"根因验证", "每个关键原因必须有验证方式和证据，不能只凭经验判断。")
    if factors:
        for i,f in enumerate(factors[:4]):
            xx=1.0+(i%2)*5.65; yy=1.78+(i//2)*1.58
            shape(s,xx,yy,5.10,1.18,"white","line")
            shape(s,xx,yy,0.48,1.18,["red","orange","blue","green"][i],["red","orange","blue","green"][i],False)
            add_text(s,str(i+1),xx+0.18,yy+0.50,0.12,0.09,9,"white",True,"center")
            add_text(s,_short(f.get("cause","关键原因"),26),xx+0.68,yy+0.18,3.95,0.16,9.2,"black",True)
            add_text(s,"验证："+_short(f.get("verification_method","待补充"),34),xx+0.68,yy+0.50,4.10,0.13,7.8,"grey")
            add_text(s,"证据："+_short(f.get("evidence","待补充"),34),xx+0.68,yy+0.78,4.10,0.13,7.8,"grey")
    else:
        rich_table(s,0.95,1.85,[2.8,2.7,3.1,2.6],0.60,["疑似根因","验证方法","证据/样本","结论"],[['待补充','待补充','待补充','待补充']],accent_cols={0},max_rows=1)
    conclusion(s,"只保留有证据支撑、能够被工程对策解决的关键原因。")


def five_w_countermeasure_page(prs, data, page):
    items=data.get("qcc_process",{}).get("countermeasures",{}).get("items",[])[:4]
    s=slide(prs,page,"拟定对策｜5W", "每项对策必须说清 What/Why/Who/When/Where，避免措施清单化。")
    if not items:
        items=[{"countermeasure":"待补充对策","key_factor":"待补充根因","owner":"待补充","deadline":"待补充","scope":"待补充范围"}]
    for i,it in enumerate(items[:4]):
        xx=1.0+(i%2)*5.65; yy=1.72+(i//2)*1.70
        shape(s,xx,yy,5.15,1.30,"white","line")
        add_text(s,f"{i+1}. What",xx+0.18,yy+0.18,0.95,0.12,8.2,"red",True)
        add_text(s,_short(it.get("countermeasure") or it.get("implementation"),32),xx+1.12,yy+0.18,3.70,0.12,8.2,"black",True)
        add_text(s,"Why："+_short(it.get("key_factor","对应要因"),27),xx+0.18,yy+0.47,4.60,0.10,8.0,"grey")
        add_text(s,"Who："+_short(it.get("owner","QCC小组"),12)+"   When："+_short(it.get("deadline","评审后"),12),xx+0.18,yy+0.72,4.60,0.10,8.0,"grey")
        add_text(s,"Where："+_short(it.get("scope") or it.get("validation") or "目标流程/系统",30),xx+0.18,yy+0.96,4.60,0.10,8.0,"grey")
    conclusion(s,"每项对策都必须对应关键要因，并绑定负责人、范围、时间和验证方式。")


def five_w_implementation_page(prs, page):
    s=slide(prs,page,"实施跟踪｜5W", "实施页保留 5W 追踪口径，确保执行过程可检查。")
    items=[
        ("闭环控制链路","建立采样、诊断、决策、执行闭环","QCC小组","本轮活动","线程池治理链路"),
        ("SafetyGate","控制误调、频繁调和不可回退风险","架构/测试","本轮活动","动态调整入口"),
        ("Evidence","形成审计、复盘和回滚证据链","开发/运维","本轮活动","调整记录与报告"),
        ("重建重放","解决队列不可热改的工程约束","开发/测试","本轮活动","Executor 生命周期"),
    ]
    headers=["What","Why","Who","When","Where"]
    rich_table(s,0.95,1.80,[2.00,3.20,1.60,1.55,2.75],0.52,headers,items,accent_cols={0},max_rows=4)
    card(s,1.0,5.38,11.30,0.66,"实施结论","执行跟踪必须能回到 5W：做什么、为什么、谁负责、何时完成、在哪个范围生效。","red")


def standardization_checklist_page(prs, data, page):
    assets=data.get("standardization",{}).get("assets",[])[:6]
    if not assets:
        assets=[{"type":"流程标准","content":"待补充"},{"type":"数据标准","content":"待补充"},{"type":"测试标准","content":"待补充"}]
    s=slide(prs,page,"成果固化｜标准化清单", "把有效做法固化为流程、数据、测试和运维资产，防止反弹。")
    for i,a in enumerate(assets[:6]):
        yy=1.70+i*0.56
        shape(s,1.00,yy,11.15,0.42,"white","line",False)
        shape(s,1.00,yy,0.42,0.42,["red","orange","blue","green","purple","red"][i%6],["red","orange","blue","green","purple","red"][i%6],False)
        add_text(s,"✓",1.14,yy+0.15,0.12,0.08,8.5,"white",True,"center")
        add_text(s,a.get("type",f"清单{i+1}"),1.62,yy+0.14,1.75,0.08,8.8,"red",True)
        add_text(s,_short(a.get("content","待补充"),68),3.50,yy+0.14,8.20,0.08,8.2,"black")
    card(s,1.0,5.52,11.30,0.55,"固化结论","标准化清单应明确固化项、规则/脚本/流程、责任人和后续检查方式。","red")

def cover(prs, data):
    s = slide(prs)
    title = data["meta"].get("title", "QCC成果汇报")
    add_text(s, title, 0.95, 1.45, 8.6, 0.55, 30, "black", True)
    line(s, 0.98, 2.12, 2.0, 2.12, "red", 2.0)
    add_text(s, data["meta"].get("subtitle", ""), 0.98, 2.35, 7.0, 0.28, 15, "grey")
    meta = [
        f"部门：{data['meta'].get('department','待补充')}",
        f"作者：{data['meta'].get('author','待补充')}",
        f"日期：{data['meta'].get('date','待补充')}",
        f"Security Level: {data['meta'].get('security_level','Huawei Confidential')}"
    ]
    for i, m in enumerate(meta): add_text(s, m, 0.98, 4.45+i*0.32, 4.5, 0.18, 11, "black")
    # Do not add a duplicate HUAWEI logo. The approved template owns logo placement.



def agenda(prs):
    s = slide(prs, 2)
    add_text(s, "目录", 0.95, 0.72, 2.2, 0.45, 28, "black", True)
    line(s, 0.98, 1.28, 1.9, 1.28, "red", 2)
    items = [
        ("01", "课题评审", "头脑风暴 / 亲和图 / 检查表"),
        ("02", "把握现状", "SIPOC / 柏拉图 / 子流程图 / 目标设定"),
        ("03", "根因与对策", "鱼骨图 / 矩阵图 / 根因验证 / 5W"),
        ("04", "验证与固化", "效果确认 / 风险关闭 / 标准化清单 / 推广")]
    for i,(num, t, d) in enumerate(items):
        y=2.0+i*0.9
        add_text(s,num,1.2,y,0.6,0.25,18,"red",True)
        add_text(s,t,2.0,y+0.02,2.4,0.20,14,"black",True)
        add_text(s,d,4.3,y+0.04,6.5,0.18,10,"grey")

def circle_profile(prs, data, page):
    c = data["circle"]
    s = slide(prs, page, "圈组介绍：以角色分工、证据责任和活动机制保障 QCC 闭环", "圈组页用于建立活动可信度，而不是简单罗列成员。")
    card(s,0.95,1.85,3.4,1.35,"圈组定位",f"圈名：{c.get('name')}\n课题：{c.get('topic')}\n周期：{c.get('activity_period')}","red")
    members = c.get("members", [])[:4]
    role_text = "\n".join([f"{m.get('role')}：{m.get('responsibility')}" for m in members])
    card(s,4.65,1.85,3.55,1.35,"角色与责任",role_text,"orange")
    scope = c.get("evidence_scope", {})
    card(s,8.5,1.85,3.4,1.35,"证据范围",f"项目：{scope.get('project')}\n分支：{scope.get('branch')}\n目录：{scope.get('source_dir')}","blue")
    acts = "\n".join([f"• {x}" for x in c.get("activity_mechanism", [])])
    card(s,1.25,4.0,10.5,1.25,"活动机制",acts,"green")



def topic_current_target(prs, data):
    brainstorm_page(prs, data, 4)
    affinity_page(prs, data, 5)
    checklist_topic_page(prs, data, 6)
    sipoc_page(prs, data, 7)
    pareto_page(prs, data, 8)
    subprocess_page(prs, data, 9)

    # Target setting
    s = slide(prs,10,"目标设定：用指标组合约束改善边界", "目标页同时表达结果指标、过程指标和安全边界。")
    tg=data["qcc_process"].get("target_setting",{}).get("targets",[])[:4]
    for i,t in enumerate(tg):
        kpi(s,1.0+i*2.95,1.85,2.3,1.12,str(t.get("target")),t.get("metric"),t.get("rationale"),["green","blue","red","orange"][i%4])
    if not tg:
        kpi(s,1.0,1.85,2.3,1.12,"待补充","目标指标","请补充目标设定","red")
    card(s,1.0,4.10,11.3,1.12,"目标边界","不是追求无限扩容，而是在安全门控内根据负载调整；不是替代运维判断，而是先进入影子模式，先建议、后执行；不是只验证 happy path，而要覆盖振荡、过载、参数敏感性和重建失败。","red")


def cause_key_countermeasure(prs, data):
    q=data["qcc_process"]
    s=slide(prs,11,"根因分析｜鱼骨图", "六类维度不是简单罗列，需要最终收敛到可行动的关键原因。")
    dims=q.get("cause_analysis",{}).get("dimensions",[])[:6]
    fishbone(s,0.75,1.35,dims,"六类原因进一步收敛为：运行感知不足、执行与调整受限、质量与管控缺失。")

    cause_matrix_page(prs, data, 12)
    root_validation_page(prs, data, 13)
    five_w_countermeasure_page(prs, data, 14)


def implementation_slides(prs):
    five_w_implementation_page(prs, 15)

    s=slide(prs,16,"对策实施一：构建采样、诊断、决策、执行的闭环控制链路", "从人工经验调参转为数据驱动、策略受控、执行可追溯的治理模式。")
    steps=[("采样","队列长度\n活跃线程\n拒绝/延迟","red"),("诊断","高负载\n下游阻塞\n低峰空闲","orange"),("决策","扩容/缩容\n队列调整\n保持观察","blue"),("安全门","冷却时间\n限额\n方向阻断","green"),("执行","调整线程池\n证据记录\n回滚预案","red")]
    flow_bar(s,0.95,2.10,steps,box_w=1.85,box_h=1.45,gap=0.43)
    card(s,1.0,4.70,11.35,0.82,"闭环价值","把分钟级人工处理链路压缩为毫秒级治理链路，同时通过安全门和证据链控制误调、频繁调和不可回退风险。","red")

    s=slide(prs,17,"对策实施二：SafetyGate 将动态调整变成可控治理", "动态机制最大的风险不是不会调，而是误调、频繁调、不可回退。")
    guards=[("01","准入检查","状态健康、指标完整"),("02","冷却控制","避免连续频繁变更"),("03","幅度限制","单次调整有上限"),("04","方向阻断","避免扩缩反复横跳"),("05","拒绝保护","拒绝风险优先处理"),("06","证据校验","执行前后均可追溯"),("07","回滚预案","异常时恢复安全配置"),("08","影子开关","先建议，后执行")]
    gate_grid(s,1.0,1.85,guards)
    card(s,1.05,5.25,11.25,0.66,"控制逻辑","先检查，再限频、限幅、限方向，同时保留证据、回滚与影子模式，确保能调、敢调、可追溯。","red")

    s=slide(prs,18,"对策实施三：Evidence 让每一次调整都可审计、可复盘、可回滚", "QCC 改善不能只看最终结果，过程证据同样需要标准化。")
    ev=[("Before","调整前快照\n线程池配置 / 队列长度 / 负载状态","red"),("Command","治理命令\ncommandId / 策略结论 / 安全门校验","orange"),("After","调整后结果\n执行结果 / 拒绝延迟 / 资源变化","blue"),("Audit","审计沉淀\nJSONL记录 / 回放报告 / 异常追溯","green")]
    flow_bar(s,1.0,2.15,ev,box_w=2.25,box_h=1.50,gap=0.58)
    card(s,1.0,4.85,11.35,0.78,"标准化要求","每次调整必须有唯一 commandId；执行前后关键指标成对记录；影子模式输出建议，不直接变更生产；失败路径必须可复盘。","red")

    s=slide(prs,19,"关键技术障碍：队列容量不可热改，通过重建与重放实现闭环", "这不是简单 setter 问题，而是 JDK 执行器模型决定的工程约束。")
    steps=[("约束","ThreadPoolExecutor 原生不支持安全替换队列；直接改队列可能破坏任务语义。","red"),("对策","隔离旧 Executor；安全创建新 Executor；未完成任务可控重放。","orange"),("结果","支持动态队列治理；避免任务丢失；重建路径可测试。","green")]
    flow_bar(s,1.05,2.25,steps,box_w=3.05,box_h=1.70,gap=0.95)
    card(s,1.0,4.88,11.35,0.72,"工程判断","复杂技术障碍页要明确表达约束、对策和结果之间的因果链路。","red")


def results_slides(prs, data):
    res=data["results"]
    s=slide(prs,20,"验证体系：用自动化用例形成质量门禁", "技术型 QCC 需要把有效变成可重复验证的测试资产。")
    v=res.get("validation",{})
    kpis=[(v.get("test_passed"),"测试通过","0失败/0错误"),(v.get("experiment_count"),"实验场景","覆盖峰谷、阻塞、过载"),("40","交替负载步数","验证无振荡"),("5×","过载压力","验证退化边界")]
    for i,k in enumerate(kpis): kpi(s,1.0+i*2.85,1.80,2.2,1.05,str(k[0]),k[1],k[2],["green","red","blue","orange"][i])
    layers=[("单元测试","锁竞争、生命周期、卡死检测"),("集成测试","Redis锁、Kafka链路、执行审计"),("E2E测试","取消、重建、异常恢复路径"),("实验报告","唯一执行、卡死恢复、风险关闭")]
    vertical_timeline(s,1.0,3.45,[(a,b,["red","orange","blue","green"][i]) for i,(a,b) in enumerate(layers)],step_h=0.62)

    s=slide(prs,21,"效果确认：用 Before/After 表格表达改善结果", "改善结果必须 before/after 对比表达，而不是只展示实验结论。")
    ba=res.get("before_after",[])[:5]
    rows=[[b.get("metric",""), b.get("before",""), b.get("after",""), b.get("improvement","")] for b in ba]
    rich_table(s,0.95,1.75,[2.55,2.65,2.45,3.55],0.56,["指标","改善前","改善后","改善幅度"],rows,accent_cols={3},max_rows=5)
    kpi(s,1.35,5.45,2.25,0.95,"32%","平均线程占用降低","8 → 5.4","orange")
    kpi(s,5.4,5.45,2.25,0.95,"0","拒绝与振荡","稳定性风险关闭","green")
    kpi(s,9.45,5.45,2.25,0.95,"≤3ms","p95差距","性能边界可控","blue")

    s=slide(prs,22,"风险关闭：把动态治理的主要担忧逐项转化为实验结论", "优秀 QCC 不回避风险，而是把每个风险假设绑定实验、结果与处置结论。")
    risks=res.get("risk_closure",[])[:6]
    rows=[[r.get("concern",""), r.get("experiment",""), r.get("result",""), r.get("conclusion","")] for r in risks]
    rich_table(s,0.95,1.70,[2.25,3.0,3.0,3.05],0.56,["担忧","实验","结果","结论"],rows,accent_cols={0},max_rows=6)


def std_promotion_summary(prs, data):
    s=slide(prs,23,"成果总结：形成可量化效果、可复用资产和可推广方法", "QCC 成果既包括指标改善，也包括团队能力和标准化资产沉淀。")
    rows=[
        ["有形成果","646个自动化测试通过；16组实验形成验证基线；0拒绝、0振荡；168ms闭环扩容响应；平均线程占用降低32%"],
        ["无形成果","建立线程池动态治理方法；降低经验调参依赖；形成风险假设到实验关闭习惯；支撑影子模式推广"],
        ["标准化成果","运行规范；测试门禁；Evidence数据规范；灰度与回滚流程；AI辅助开发报告模板"],
    ]
    rich_table(s,1.0,1.85,[2.2,9.2],0.72,["成果类型","成果内容"],rows,accent_cols={0},max_rows=3)

    standardization_checklist_page(prs, data, 24)

    s=slide(prs,25,"推广计划：从 Demo 验证进入影子模式，再逐步灰度到生产服务", "动态线程池治理必须遵循先观察、再建议、后执行的推广节奏。")
    roads=data["promotion"].get("roadmap",[])[:3]
    timeline=[]
    for i,r in enumerate(roads):
        body="；".join(r.get("content",[]))
        timeline.append((r.get("phase",""),body,["red","orange","blue"][i%3]))
    if not timeline:
        timeline=[("短期","补齐数据与影子模式报告","red"),("中期","选择低风险服务试点","orange"),("长期","标准化纳管与推广","blue")]
    vertical_timeline(s,1.0,1.95,timeline,step_h=1.10)
    card(s,1.05,5.55,11.25,0.55,"推广原则","先观察、再建议、后执行；先单点影子验证，再灰度扩展，最后进入标准化纳管。","red")

    s=slide(prs,26,"总结：本轮 QCC 完成了从经验调参到闭环治理的质量改进", "通过完整 PDCA 闭环，形成了可验证、可审计、可推广的治理方法。")
    pdca_cycle(s,3.1,1.70)
    card(s,1.05,5.00,11.2,0.88,"核心结论","动态治理的关键不是自动调大线程数，而是建立有目标、有要因、有对策、有证据、有标准的运行质量治理系统。","red")

    # End page is template-owned. Do not add another Thank-you, copyright,
    # vision statement, legal block, or HUAWEI logo on top of the supplied end
    # layout. The template provides the complete ending page.
    slide(prs,27, role="end")

def build(data: dict, template: Path | None) -> Presentation:
    if not template or not template.exists():
        raise FileNotFoundError("v2.4 requires an approved template file; blank Presentation() fallback is forbidden")
    prs = Presentation(str(template))
    configure_template_layouts(prs)
    sanitize_template_layouts()
    remove_template_seed_slides(prs)
    prs.slide_width = Inches(W_IN); prs.slide_height = Inches(H_IN)
    cover(prs, data); agenda(prs); circle_profile(prs, data, 3); topic_current_target(prs, data); cause_key_countermeasure(prs, data); implementation_slides(prs); results_slides(prs, data); std_promotion_summary(prs, data)
    return prs


def main():
    global FOOTER_MODE
    ap=argparse.ArgumentParser()
    ap.add_argument("qcc_data", type=Path, nargs="?", default=None, help="Optional. If omitted, generate the bundled empty QCC template instead of inventing content.")
    ap.add_argument("--template", type=Path, default=None, help="Optional custom template. If omitted, use templates/template_light_16_9.pptx bundled with this Skill.")
    ap.add_argument("--empty-template", action="store_true", help="Force empty QCC template output, ignoring qcc_data.")
    ap.add_argument("--output", type=Path, required=True)
    ap.add_argument("--footer-mode", choices=["none", "minimal", "full"], default="none",
                    help="Default is none because the approved template owns footer/logo. full is disabled in v2.4.")
    args=ap.parse_args()
    FOOTER_MODE = args.footer_mode
    skill_root = Path(__file__).resolve().parents[1]
    default_template = skill_root / "templates" / "template_light_16_9.pptx"
    empty_template = skill_root / "templates" / "qcc-empty-template.pptx"
    args.output.parent.mkdir(parents=True, exist_ok=True)

    if args.empty_template or args.qcc_data is None:
        if not empty_template.exists():
            raise FileNotFoundError(f"empty template not found: {empty_template}")
        shutil.copyfile(empty_template, args.output)
        print(f"written empty QCC template: {args.output}")
        return

    template = args.template or default_template
    data=load_yaml(args.qcc_data)
    prs=build(data,template)
    prs.save(args.output)
    print(f"written: {args.output}")

if __name__ == "__main__":
    main()
