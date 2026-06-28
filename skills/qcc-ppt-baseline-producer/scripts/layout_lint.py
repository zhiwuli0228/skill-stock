#!/usr/bin/env python3
"""Static layout lint for generated PPTX.

Formal mode treats true layout warnings as failures. Heuristic whitespace findings
are reported as notes because v2.3 removes template placeholders that previously
inflated content bounds.
"""
from __future__ import annotations
import argparse
from pathlib import Path
from pptx import Presentation

EMU_PER_INCH = 914400
BOTTOM_SAFE_Y = 6.72
LOGO_SAFE_X = 10.80
LOGO_SAFE_Y = 6.72


def inch(v): return v / EMU_PER_INCH


def shape_bounds(shp):
    return inch(shp.left), inch(shp.top), inch(shp.width), inch(shp.height)


def has_text(shp):
    return hasattr(shp, "text") and bool((shp.text or "").strip())


def min_font_size(shp):
    vals=[]
    if not getattr(shp, "has_text_frame", False): return None
    for p in shp.text_frame.paragraphs:
        for r in p.runs:
            if r.font.size:
                vals.append(r.font.size.pt)
    return min(vals) if vals else None


def is_bottom_conclusion(shp):
    if not has_text(shp): return False
    x,y,w,h = shape_bounds(shp)
    text = shp.text.strip()
    return y > 5.85 and h < 0.50 and any(k in text for k in ["本页结论", "风险关闭", "闭环价值", "控制逻辑", "验证逻辑", "核心结论", "结论"])


def in_logo_safe_zone(shp):
    x,y,w,h = shape_bounds(shp)
    return x + w > LOGO_SAFE_X and y + h > LOGO_SAFE_Y


def is_generated_huawei_text(shp):
    if not has_text(shp):
        return False
    normalized = " ".join(shp.text.upper().split())
    return normalized == "HUAWEI"


def text_len(shp):
    return len((getattr(shp, "text", "") or "").strip())


def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("pptx", type=Path)
    ap.add_argument("--min-body-font", type=float, default=7.8)
    ap.add_argument("--strict", action="store_true", help="Fail on warnings. Enabled automatically in formal mode.")
    ap.add_argument("--mode", choices=["formal", "draft"], default="formal")
    args=ap.parse_args()

    prs=Presentation(str(args.pptx))
    sw, sh = inch(prs.slide_width), inch(prs.slide_height)
    errors=[]; warnings=[]; notes=[]

    for idx, slide in enumerate(prs.slides, 1):
        layout_name = (slide.slide_layout.name or "").lower()
        if idx == 1 and not any(k in layout_name for k in ["探索", "cover", "title"]):
            errors.append(f"slide {idx}: cover must use the approved cover/template-title layout, got '{slide.slide_layout.name}'")
        if idx == 2 and not any(k in layout_name for k in ["contents", "目录", "agenda"]):
            errors.append(f"slide {idx}: agenda must use the approved contents/agenda layout, got '{slide.slide_layout.name}'")
        if idx not in (1, 2, len(prs.slides)) and any(k in layout_name for k in ["end", "结束"]):
            errors.append(f"slide {idx}: content slide is using end-page layout '{slide.slide_layout.name}'")
        if idx == len(prs.slides) and not any(k in layout_name for k in ["end", "结束"]):
            warnings.append(f"slide {idx}: final slide should use the approved end-page layout, got '{slide.slide_layout.name}'")
        bottom_bars=0
        max_bottom=0.0
        dense_text_shapes=0
        huawei_texts=0
        for j, shp in enumerate(slide.shapes):
            try: x,y,w,h=shape_bounds(shp)
            except Exception: continue
            if has_text(shp) and (w <= 0.01 or h <= 0.01):
                errors.append(f"slide {idx}: text shape {j} has non-positive/invalid size x={x:.2f},y={y:.2f},w={w:.2f},h={h:.2f}: {getattr(shp,'text','')[:30]}")
            if x < -0.01 or y < -0.01 or x+w > sw+0.01 or y+h > sh+0.01:
                errors.append(f"slide {idx}: shape {j} out of bounds x={x:.2f},y={y:.2f},w={w:.2f},h={h:.2f}")
            if y+h > max_bottom and y < BOTTOM_SAFE_Y:
                max_bottom=y+h
            fs=min_font_size(shp)
            if fs is not None and fs < args.min_body_font:
                warnings.append(f"slide {idx}: tiny font {fs:.1f}pt in shape {j}: {getattr(shp,'text','')[:36]}")
            if is_bottom_conclusion(shp):
                bottom_bars += 1
            if has_text(shp) and text_len(shp) > 18:
                dense_text_shapes += 1
            if is_generated_huawei_text(shp):
                huawei_texts += 1
            # Detect compact cards rendered with title/body text boxes that cannot
            # physically fit. This catches the common failure where text is visible
            # outside a very short row even though the slide canvas bounds are valid.
            if has_text(shp) and y > 3.2 and h < 0.08:
                errors.append(f"slide {idx}: text box height too small in shape {j}: y={y:.2f},h={h:.2f}, text={shp.text[:30]}")
            if in_logo_safe_zone(shp) and has_text(shp) and idx not in (1, len(prs.slides)):
                # Footer text near the reserved template/logo area is a common cause of overlap.
                txt = " ".join(shp.text.strip().split())
                if txt.upper() == "HUAWEI" or (len(txt) > 8 and "CONFIDENTIAL" not in txt.upper()):
                    errors.append(f"slide {idx}: text occupies bottom-right template/logo safe zone in shape {j}: {shp.text[:30]}")
        if huawei_texts:
            errors.append(f"slide {idx}: generated HUAWEI text detected; do not add logo text on top of the supplied template")
        if bottom_bars > 1:
            errors.append(f"slide {idx}: more than one bottom conclusion-like text ({bottom_bars})")
        if idx in (1,2,3) and bottom_bars > 0:
            errors.append(f"slide {idx}: conclusion bar not allowed on cover/agenda/circle page")
        # Warn if dense content stops too early: cover/agenda/thank-you excluded.
        if idx not in (1,2,len(prs.slides)) and dense_text_shapes >= 3 and max_bottom < 5.65:
            # After v2.3 placeholder cleanup, this heuristic can produce false positives
            # because unused template placeholders no longer inflate the content bottom.
            # Keep it as a non-blocking note; visual inspection remains mandatory.
            notes.append(f"slide {idx}: possible unused vertical space; content bottom at {max_bottom:.2f}in")

    if errors:
        print("ERRORS:")
        for e in errors: print("-", e)
    if warnings:
        print("WARNINGS:")
        for w in warnings: print("-", w)
    if notes:
        print("NOTES:")
        for n in notes: print("-", n)
    fail_on_warnings = args.strict or args.mode == "formal"
    if errors or (fail_on_warnings and warnings):
        print(f"FAILED: {len(errors)} errors, {len(warnings)} warnings")
        return 1
    print(f"PASSED: 0 errors, {len(warnings)} warnings, {len(notes)} notes")
    return 0

if __name__ == "__main__":
    raise SystemExit(main())
