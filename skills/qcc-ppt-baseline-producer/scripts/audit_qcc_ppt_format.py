#!/usr/bin/env python3
"""Audit basic format risks in a QCC PowerPoint deck.

This is intentionally conservative: it does not replace rendered visual review,
but it catches repeatable anti-patterns such as dense tables, overloaded slides,
small body fonts, long titles, out-of-bounds shapes, and shape-count explosions.
"""
from __future__ import annotations

import argparse
from pathlib import Path
from pptx import Presentation

EMU_PER_INCH = 914400

DEFAULT_BODY_MIN_FONT = 9.0
DEFAULT_MAX_TEXTBOXES = 36
DEFAULT_MAX_CHARS = 620
DEFAULT_MAX_SHAPES = 72
DEFAULT_MAX_TABLE_ROWS = 6
DEFAULT_MAX_TABLE_COLS = 6
DEFAULT_MAX_TITLE_CHARS = 34


def emu_to_in(v: int) -> float:
    return float(v) / EMU_PER_INCH


def is_footer_text(txt: str) -> bool:
    t = txt.strip().lower()
    return t.isdigit() or 'confidential' in t or 'huawei' in t or 'copyright' in t


def iter_text_runs(shape):
    if not getattr(shape, 'has_text_frame', False):
        return
    txt = shape.text.strip()
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            if r.font.size is not None:
                yield txt, r.font.size.pt


def table_shape_info(shape):
    if not getattr(shape, 'has_table', False):
        return None
    table = shape.table
    return len(table.rows), len(table.columns)


def audit(pptx: Path) -> tuple[list[dict], dict]:
    prs = Presentation(str(pptx))
    width = prs.slide_width
    height = prs.slide_height
    issues = []
    summary = {'slides': len(prs.slides), 'issue_count': 0}
    for idx, slide in enumerate(prs.slides, start=1):
        textboxes = 0
        chars = 0
        shapes = len(slide.shapes)
        min_font = None
        title_text = None
        small_font_count = 0
        table_risks = []
        oob = []
        for shape in slide.shapes:
            # out-of-bounds check
            try:
                if shape.left < 0 or shape.top < 0 or shape.left + shape.width > width or shape.top + shape.height > height:
                    oob.append(getattr(shape, 'name', 'shape'))
            except Exception:
                pass

            if getattr(shape, 'has_text_frame', False):
                txt = shape.text.strip()
                if txt:
                    textboxes += 1
                    chars += len(txt)
                    if title_text is None and len(txt) > 3 and not is_footer_text(txt):
                        title_text = txt.replace('\n', ' ')
                    for run_txt, pt in iter_text_runs(shape) or []:
                        if is_footer_text(run_txt):
                            continue
                        min_font = pt if min_font is None else min(min_font, pt)
                        if pt < DEFAULT_BODY_MIN_FONT:
                            small_font_count += 1
            info = table_shape_info(shape)
            if info:
                rows, cols = info
                if rows > DEFAULT_MAX_TABLE_ROWS or cols > DEFAULT_MAX_TABLE_COLS:
                    table_risks.append(f'{rows}x{cols}')

        slide_issues = []
        if title_text and len(title_text) > DEFAULT_MAX_TITLE_CHARS:
            slide_issues.append(f'long title ({len(title_text)} chars)')
        if textboxes > DEFAULT_MAX_TEXTBOXES:
            slide_issues.append(f'too many text boxes ({textboxes})')
        if chars > DEFAULT_MAX_CHARS:
            slide_issues.append(f'text overload ({chars} chars)')
        if shapes > DEFAULT_MAX_SHAPES:
            slide_issues.append(f'too many shapes ({shapes})')
        if min_font is not None and min_font < DEFAULT_BODY_MIN_FONT:
            slide_issues.append(f'small body font min={min_font:.1f}pt, runs={small_font_count}')
        if table_risks:
            slide_issues.append(f'dense table(s): {", ".join(table_risks)}')
        if oob:
            slide_issues.append(f'out-of-bounds shape(s): {len(oob)}')
        if slide_issues:
            issues.append({
                'slide': idx,
                'title': title_text or '(no title)',
                'issues': slide_issues,
                'metrics': {'textboxes': textboxes, 'chars': chars, 'shapes': shapes, 'min_font': min_font},
            })
    summary['issue_count'] = len(issues)
    return issues, summary


def write_report(pptx: Path, report: Path, issues: list[dict], summary: dict) -> None:
    lines = []
    lines.append('# QCC PPT Format Audit Report')
    lines.append('')
    lines.append(f'- PPTX: `{pptx}`')
    lines.append(f'- Slides: {summary["slides"]}')
    lines.append(f'- Slides with format risk: {summary["issue_count"]}')
    lines.append('')
    if not issues:
        lines.append('## Result')
        lines.append('')
        lines.append('PASS: no obvious programmable format risks found. Render screenshots still require visual inspection.')
    else:
        lines.append('## Risks')
        lines.append('')
        lines.append('| Slide | Title | Risks |')
        lines.append('|---:|---|---|')
        for item in issues:
            risks = '<br>'.join(item['issues'])
            title = item['title'].replace('|', '/').replace('\n', ' ')[:80]
            lines.append(f'| {item["slide"]} | {title} | {risks} |')
        lines.append('')
        lines.append('## Interpretation')
        lines.append('')
        lines.append('- This report flags mechanical risks only. A slide can be acceptable after rendered review even if flagged.')
        lines.append('- Dense QCC method tables should be split, cardized, or reduced to key rows before final delivery.')
        lines.append('- Body text below 9pt, excluding footer/legal pages, should be treated as a formal review risk.')
    report.parent.mkdir(parents=True, exist_ok=True)
    report.write_text('\n'.join(lines), encoding='utf-8')


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('pptx', type=Path)
    parser.add_argument('--report', type=Path, default=Path('qcc-format-audit-report.md'))
    args = parser.parse_args()
    issues, summary = audit(args.pptx)
    write_report(args.pptx, args.report, issues, summary)
    print(f'Wrote {args.report}; slides_with_risk={summary["issue_count"]}')


if __name__ == '__main__':
    main()
