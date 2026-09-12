#!/usr/bin/env python3
"""Self-test for the standalone visual enhancer package.

Builds a minimal sample deck (with a circular badge), runs the deterministic entry
point, and asserts that:

* the enhanced deck exists and keeps the same slide count;
* the circular badge was centered;
* the visual review report was written.
"""
from __future__ import annotations

import argparse
import subprocess
import sys
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


def build_sample(path: Path) -> int:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    title = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(8), Inches(0.6))
    title.text_frame.text = "样例基线：徽标居中自检"
    run = title.text_frame.paragraphs[0].runs[0]
    run.font.size = Pt(22)
    run.font.bold = True

    badge = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(1.0), Inches(2.0), Inches(0.6), Inches(0.6)
    )
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(0xC8, 0x10, 0x2E)
    badge.line.fill.background()

    label = slide.shapes.add_textbox(
        Inches(1.06), Inches(2.06), Inches(0.5), Inches(0.5)
    )
    label.text_frame.text = "1"
    label_run = label.text_frame.paragraphs[0].runs[0]
    label_run.font.size = Pt(16)
    label_run.font.bold = True
    label_run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    prs.save(str(path))
    return len(prs.slides._sldIdLst)


def main() -> int:
    parser = argparse.ArgumentParser(description="Self-test the enhancer package.")
    parser.add_argument("--keep", action="store_true", help="keep the temp workspace")
    args = parser.parse_args()

    skill_dir = Path(__file__).resolve().parent.parent
    entry = skill_dir / "scripts" / "enhance_qcc_ppt.py"
    failures: list[str] = []

    with tempfile.TemporaryDirectory(prefix="qcc-enhancer-") as tmp:
        workspace = Path(tmp)
        sample = workspace / "input" / "qcc-baseline.pptx"
        sample.parent.mkdir(parents=True, exist_ok=True)
        slides = build_sample(sample)

        output = workspace / "output" / "qcc-review-ready.pptx"
        report = workspace / "reports" / "visual-review-report.md"
        completed = subprocess.run(
            [
                sys.executable,
                str(entry),
                "--input",
                str(sample),
                "--output",
                str(output),
                "--report",
                str(report),
            ],
            capture_output=True,
            text=True,
            encoding="utf-8",
        )
        print(completed.stdout.strip())
        if completed.stderr.strip():
            print(completed.stderr.strip())

        if completed.returncode != 0:
            failures.append("enhance_qcc_ppt.py returned a non-zero exit code")
        if not output.exists():
            failures.append("enhanced deck was not produced")
        else:
            produced = Presentation(str(output))
            if len(produced.slides._sldIdLst) != slides:
                failures.append("slide count changed during enhancement")
        if "centered_badges=0" in completed.stdout:
            failures.append("circular badge was not centered")
        if not report.exists():
            failures.append("visual review report was not written")

        if args.keep:
            target = Path.cwd() / "selftest-workspace"
            if target.exists():
                failures.append(f"refusing to overwrite existing {target}")
            else:
                import shutil

                shutil.copytree(workspace, target)
                print(f"kept workspace: {target}")

    if failures:
        print("SELFTEST FAILED:")
        for item in failures:
            print(f"- {item}")
        return 1
    print("SELFTEST PASSED: enhancer package runs standalone.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
