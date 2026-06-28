#!/usr/bin/env python3
"""Center numeric/letter badge text in circular PowerPoint shapes.

This post-build repair is intentionally conservative:
- it finds oval shapes;
- finds nearby one/two-character numeric or uppercase badge text;
- resizes the badge text box to the oval bounds;
- sets zero margins, center alignment, and middle vertical anchor.

Usage:
    python scripts/center_badge_text.py input.pptx --output output.pptx
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path
from collections import Counter

from pptx import Presentation
from pptx.enum.text import MSO_VERTICAL_ANCHOR, PP_ALIGN

EMU_PER_INCH = 914400
BADGE_RE = re.compile(r"^[A-Z]{1,2}$|^[0-9]{1,2}$")
EXEMPT_TEXT = {"AI"}


def is_oval(shape) -> bool:
    try:
        return "OVAL" in str(shape.auto_shape_type)
    except Exception:
        return False


def is_badge_text(text: str) -> bool:
    value = text.strip()
    return bool(BADGE_RE.fullmatch(value)) and value not in EXEMPT_TEXT


def center_badges(pptx_path: Path, output_path: Path) -> Counter:
    prs = Presentation(str(pptx_path))
    changes: Counter = Counter()

    for slide_idx, slide in enumerate(prs.slides, start=1):
        ovals = [shape for shape in slide.shapes if is_oval(shape)]
        candidates = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                if is_badge_text(text):
                    candidates.append(shape)

        used = set()
        for oval in ovals:
            ocx = oval.left + oval.width / 2
            ocy = oval.top + oval.height / 2
            best = None
            best_dist = float("inf")
            for text_shape in candidates:
                if id(text_shape) in used:
                    continue
                tcx = text_shape.left + text_shape.width / 2
                tcy = text_shape.top + text_shape.height / 2
                dx = (tcx - ocx) / EMU_PER_INCH
                dy = (tcy - ocy) / EMU_PER_INCH
                dist = dx * dx + dy * dy
                # The acceptable search radius scales with circle size.
                limit = (max(oval.width, oval.height) / EMU_PER_INCH * 0.80) ** 2
                expanded = 0.09 * EMU_PER_INCH
                center_inside_expanded_oval = (
                    oval.left - expanded <= tcx <= oval.left + oval.width + expanded
                    and oval.top - expanded <= tcy <= oval.top + oval.height + expanded
                )
                if (dist <= limit or center_inside_expanded_oval) and dist < best_dist:
                    best = text_shape
                    best_dist = dist
            if best is None:
                continue

            best.left = oval.left
            best.top = oval.top
            best.width = oval.width
            best.height = oval.height
            tf = best.text_frame
            tf.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE
            tf.margin_left = 0
            tf.margin_right = 0
            tf.margin_top = 0
            tf.margin_bottom = 0
            for paragraph in tf.paragraphs:
                paragraph.alignment = PP_ALIGN.CENTER
                paragraph.space_before = 0
                paragraph.space_after = 0
            used.add(id(best))
            changes[slide_idx] += 1

    prs.save(str(output_path))
    return changes


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--output", "-o", type=Path, required=True)
    args = parser.parse_args()

    changes = center_badges(args.pptx, args.output)
    total = sum(changes.values())
    print(f"centered_badges={total}")
    for slide_idx, count in sorted(changes.items()):
        print(f"slide {slide_idx}: {count}")


if __name__ == "__main__":
    main()
