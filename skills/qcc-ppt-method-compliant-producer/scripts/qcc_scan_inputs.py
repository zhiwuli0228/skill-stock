#!/usr/bin/env python3
"""Scan a folder of the user's own files and turn it into QCC intake evidence.

This is the default intake path: the user points at a folder that already holds
their records (Excel / Word / PowerPoint / PDF / CSV / text exports) instead of
answering a questionnaire or hand-writing YAML. The script:

1. inventories every file (type, size, hash);
2. extracts tables and text from xlsx / csv / docx / pptx / pdf / text;
3. matches the extracted blocks against the required QCC data fields;
4. writes a machine-readable evidence index (for the model to read), a human
   scan report, and a draft minimal data file that contains *only* values that
   were actually found in the evidence.

Nothing is invented: a field is prefilled only when a source block supports it.
Every other field stays ``null`` so the validator reports it as pending.

    python scripts/qcc_scan_inputs.py --dir C:/path/to/raw --workspace qcc-workspace
    python scripts/qcc_scan_inputs.py --dir C:/path/to/raw --workspace qcc-workspace --json

Exit codes: 0 = usable evidence found, 1 = folder readable but nothing usable,
            2 = bad input.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Iterable

try:  # PyYAML is a hard requirement of this Skill, but keep the import soft.
    import yaml
except ImportError:  # pragma: no cover
    yaml = None  # type: ignore[assignment]


TEXT_EXTS = {".txt", ".md", ".markdown", ".log", ".json", ".yaml", ".yml"}
TABLE_EXTS = {".csv", ".tsv"}
EXCEL_EXTS = {".xlsx", ".xlsm"}
WORD_EXTS = {".docx"}
DECK_EXTS = {".pptx"}
PDF_EXTS = {".pdf"}
LEGACY_EXTS = {".xls", ".doc", ".ppt"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".webp", ".tif", ".tiff", ".heic"}
SUPPORTED = TEXT_EXTS | TABLE_EXTS | EXCEL_EXTS | WORD_EXTS | DECK_EXTS | PDF_EXTS
SKIP_DIRS = {
    ".git", "__pycache__", "node_modules", ".venv", "venv", ".idea",
    ".vscode", ".mypy_cache", ".pytest_cache", ".svn",
}
MAX_ROWS = 200
MAX_TEXT = 8000

UNITS = "例件个次人数天周月年元块分钟小时台条件项万亿千百％%"
UNIT_RE = re.compile(f"[{re.escape(UNITS)}]")

NAME_KEYS = (
    "类别", "项目", "原因", "异常", "缺陷", "不良", "类型", "分类", "问题",
    "名称", "项次", "缺陷名称", "异常类型", "不良项目",
)
COUNT_KEYS = (
    "频次", "频数", "次数", "数量", "例数", "件数", "个数", "合计", "小计",
    "计数", "发生数", "异常数",
)
STRATA_KEYS = (
    "班别", "班次", "班组", "线别", "线体", "区域", "工段", "设备", "机台",
    "层别", "分层", "岗位", "部门", "人员", "渠道",
)
DATE_KEYS = ("期间", "周期", "日期", "时间", "月份")
SCORE_KEYS = (
    "上级政策", "重要性", "迫切性", "可行性", "圈能力", "效益性", "经济性",
    "成本", "风险", "时效性", "预期效果",
)
MEASURE_KEYS = ("对策", "措施", "改善措施", "方案")
CAUSE_KEYS = ("真因", "要因", "根本原因", "原因", "根因")
VERIFY_KEYS = ("验证", "数据来源", "结论", "验证结果", "佐证", "证据")
OWNER_KEYS = ("责任人", "负责人", "担当", "执行人", "圈员")
WHEN_KEYS = ("时间", "完成时间", "周次", "期限", "进度")
WHERE_KEYS = ("地点", "环节", "场所", "范围", "位置")
HOW_KEYS = ("做法", "如何", "措施内容", "how", "实施内容", "对策内容")


def normalise(cell: Any) -> str:
    if cell is None:
        return ""
    if isinstance(cell, bool):
        return "是" if cell else "否"
    if isinstance(cell, float) and cell.is_integer():
        return str(int(cell))
    return str(cell).strip()


def to_number(value: Any) -> float | None:
    """Parse a cell that holds a plain number, tolerating units and separators."""
    text = normalise(value).replace(",", "").replace("，", "").strip()
    if not text:
        return None
    cleaned = UNIT_RE.sub("", text).strip()
    if not re.fullmatch(r"-?\d+(?:\.\d+)?", cleaned):
        return None
    return float(cleaned)


@dataclass
class Block:
    source: str
    locator: str
    kind: str  # "table" | "text"
    rows: list[list[str]] = field(default_factory=list)
    text: str = ""
    note: str = ""

    @property
    def context(self) -> str:
        return f"{self.source} {self.locator}"

    def to_json(self) -> dict:
        payload: dict[str, Any] = {
            "source": self.source,
            "locator": self.locator,
            "kind": self.kind,
        }
        if self.note:
            payload["note"] = self.note
        if self.kind == "table":
            payload["rows"] = self.rows
        else:
            payload["text"] = self.text
        return payload


# --------------------------------------------------------------------------- #
# extraction
# --------------------------------------------------------------------------- #
def extract_csv(path: Path, rel: str) -> list[Block]:
    import csv

    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    try:
        with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
            rows = [[normalise(item) for item in row] for row in csv.reader(handle, delimiter=delimiter)]
    except OSError as error:
        return [Block(rel, "table", "text", note=f"读取失败：{error}")]
    rows = [row for row in rows if any(row)]
    if not rows:
        return []
    return [Block(rel, "table 1", "table", rows=rows[:MAX_ROWS])]


def extract_excel(path: Path, rel: str) -> list[Block]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return [Block(rel, "workbook", "text", note="未安装 openpyxl（pip install openpyxl）")]
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as error:  # noqa: BLE001
        return [Block(rel, "workbook", "text", note=f"读取失败：{error}")]
    blocks: list[Block] = []
    try:
        for sheet in workbook.worksheets:
            rows: list[list[str]] = []
            for row in sheet.iter_rows(values_only=True):
                cells = [normalise(item) for item in row]
                while cells and not cells[-1]:
                    cells.pop()
                if any(cells):
                    rows.append(cells)
                if len(rows) >= MAX_ROWS:
                    break
            if rows:
                blocks.append(Block(rel, f"sheet:{sheet.title}", "table", rows=rows))
    finally:
        workbook.close()
    return blocks


def extract_docx(path: Path, rel: str) -> list[Block]:
    try:
        import docx
    except ImportError:
        return [Block(rel, "document", "text", note="未安装 python-docx（pip install python-docx）")]
    try:
        document = docx.Document(str(path))
    except Exception as error:  # noqa: BLE001
        return [Block(rel, "document", "text", note=f"读取失败：{error}")]
    blocks: list[Block] = []
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    if paragraphs:
        blocks.append(Block(rel, "paragraphs", "text", text="\n".join(paragraphs)[:MAX_TEXT]))
    for index, table in enumerate(document.tables, 1):
        rows = [[normalise(cell.text) for cell in row.cells] for row in table.rows]
        rows = [row for row in rows if any(row)]
        if rows:
            blocks.append(Block(rel, f"table {index}", "table", rows=rows[:MAX_ROWS]))
    return blocks


def extract_pptx(path: Path, rel: str) -> list[Block]:
    try:
        from pptx import Presentation
    except ImportError:
        return [Block(rel, "deck", "text", note="未安装 python-pptx（pip install python-pptx）")]
    try:
        deck = Presentation(str(path))
    except Exception as error:  # noqa: BLE001
        return [Block(rel, "deck", "text", note=f"读取失败：{error}")]
    blocks: list[Block] = []
    for index, slide in enumerate(deck.slides, 1):
        texts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [[normalise(cell.text) for cell in row.cells] for row in shape.table.rows]
                rows = [row for row in rows if any(row)]
                if rows:
                    blocks.append(Block(rel, f"slide {index} table", "table", rows=rows[:MAX_ROWS]))
        if texts:
            blocks.append(Block(rel, f"slide {index}", "text", text="\n".join(texts)[:MAX_TEXT]))
    return blocks


def extract_pdf(path: Path, rel: str) -> list[Block]:
    try:
        import pdfplumber
    except ImportError:
        return [Block(rel, "pdf", "text", note="未安装 pdfplumber（pip install pdfplumber）")]
    blocks: list[Block] = []
    try:
        with pdfplumber.open(str(path)) as pdf:
            for index, page in enumerate(pdf.pages, 1):
                text = (page.extract_text() or "").strip()
                if text:
                    blocks.append(Block(rel, f"page {index}", "text", text=text[:MAX_TEXT]))
                if index >= 20:
                    break
    except Exception as error:  # noqa: BLE001
        return [Block(rel, "pdf", "text", note=f"读取失败：{error}")]
    return blocks


def extract_text(path: Path, rel: str) -> list[Block]:
    try:
        text = path.read_text(encoding="utf-8", errors="replace").strip()
    except OSError as error:
        return [Block(rel, "text", "text", note=f"读取失败：{error}")]
    if not text:
        return []
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if len(lines) >= 3 and all(("," in line or "\t" in line) for line in lines[:3]):
        delimiter = "\t" if "\t" in lines[0] else ","
        rows = [[item.strip() for item in line.split(delimiter)] for line in lines[:MAX_ROWS]]
        return [Block(rel, "table-like text", "table", rows=rows)]
    return [Block(rel, "text", "text", text=text[:MAX_TEXT])]


def extract_file(path: Path, rel: str) -> list[Block]:
    suffix = path.suffix.lower()
    if suffix in TABLE_EXTS:
        return extract_csv(path, rel)
    if suffix in EXCEL_EXTS:
        return extract_excel(path, rel)
    if suffix in WORD_EXTS:
        return extract_docx(path, rel)
    if suffix in DECK_EXTS:
        return extract_pptx(path, rel)
    if suffix in PDF_EXTS:
        return extract_pdf(path, rel)
    if suffix in TEXT_EXTS:
        return extract_text(path, rel)
    if suffix in LEGACY_EXTS:
        return [
            Block(rel, suffix, "text", note=f"旧格式 {suffix}：请另存为 {suffix}x 后重新扫描")
        ]
    if suffix in IMAGE_EXTS:
        return [Block(rel, suffix, "text", note="图片文件：需要视觉识别，请在报告里按文件名引用")]
    return [Block(rel, suffix or "(无扩展名)", "text", note="未支持的文件类型，仅登记清单")]


# --------------------------------------------------------------------------- #
# detection helpers
# --------------------------------------------------------------------------- #
def header_map(
    rows: list[list[str]], keys: Iterable[str], limit: int = 4
) -> tuple[int, dict[str, int]] | None:
    for index, row in enumerate(rows[:limit]):
        mapping: dict[str, int] = {}
        for column, item in enumerate(row):
            for key in keys:
                if key and key in item and key not in mapping:
                    mapping[key] = column
        if mapping:
            return index, mapping
    return None


def first_key(mapping: dict[str, int], keys: Iterable[str]) -> tuple[str, int] | None:
    for key in keys:
        if key in mapping:
            return key, mapping[key]
    return None


def cell(row: list[str], column: int | None) -> str:
    if column is None or column < 0 or column >= len(row):
        return ""
    return row[column]


def is_total(name: str) -> bool:
    return any(word in name for word in ("合计", "总计", "小计", "总数", "平均"))


def numeric_columns(rows: list[list[str]]) -> list[int]:
    if not rows:
        return []
    width = max(len(row) for row in rows)
    counts = [0] * width
    for row in rows:
        for index in range(width):
            if to_number(cell(row, index)) is not None:
                counts[index] += 1
    threshold = max(2, len(rows) // 2)
    return [index for index, value in enumerate(counts) if value >= threshold]


def keyvalue_pairs(rows: list[list[str]]) -> dict[str, str]:
    pairs: dict[str, str] = {}
    for row in rows:
        filled = [item for item in row if item]
        if len(filled) >= 2 and len(filled[0]) <= 18:
            pairs.setdefault(filled[0].rstrip("：:"), filled[1])
        for item in row:
            if "：" in item:
                key, value = item.split("：", 1)
                key, value = key.strip(), value.strip()
                if key and value and len(key) <= 18:
                    pairs.setdefault(key, value)
    return pairs


META_STOPWORDS = (
    "课题", "主题", "圈组", "圈名", "品管圈", "小组", "圈长", "组长", "辅导员",
    "成员", "指标", "衡量指标", "评价维度", "维度", "活动周期", "活动期间",
    "活动时间", "改善期间", "期间", "时间", "部门", "单位",
)


def trim_meta(value: str) -> str:
    """Cut a layout-merged line (PDF text often glues a caption onto the value)."""
    cut = len(value)
    for word in META_STOPWORDS:
        index = value.find(word, 1)
        if index > 0:
            cut = min(cut, index)
    for mark in ("（", "("):
        index = value.find(mark, 1)
        if index > 0:
            cut = min(cut, index)
    return re.sub(r"[\s｜|/、]+$", "", value[:cut]).strip()


# --------------------------------------------------------------------------- #
# detectors
# --------------------------------------------------------------------------- #
Candidate = tuple[str, Any, str, str]  # (target path, value, confidence, evidence)


def detect_categories(block: Block) -> list[Candidate]:
    rows = block.rows
    if len(rows) < 3:
        return []
    out: list[Candidate] = []
    header = header_map(
        rows,
        NAME_KEYS + COUNT_KEYS + DATE_KEYS + ("判定标准", "样本量", OWNER_KEYS[0]),
    )
    if header:
        index, mapping = header
        name_hit = first_key(mapping, NAME_KEYS)
        count_hit = first_key(mapping, COUNT_KEYS)
        same_column = bool(name_hit and count_hit and name_hit[1] == count_hit[1])
        strata_column = bool(name_hit and any(key in mapping for key in STRATA_KEYS))
        if name_hit and count_hit and not same_column and not strata_column:
            items: list[dict[str, Any]] = []
            numeric_names = 0
            for row in rows[index + 1 :]:
                name = cell(row, name_hit[1])
                count = to_number(cell(row, count_hit[1]))
                if not name or count is None or is_total(name):
                    continue
                if to_number(name) is not None:
                    numeric_names += 1
                    continue
                items.append({"name": name, "count": int(count) if float(count).is_integer() else count})
            if len(items) >= 3 and numeric_names < len(items):
                evidence = f"{block.context}（表头第 {index + 1} 行：{name_hit[0]} + {count_hit[0]}）"
                out.append(("current.categories", items, "high", evidence))
                for key, target in (
                    ("判定标准", "current.check_sheet.criteria"),
                    ("责任人", "current.check_sheet.owner"),
                ):
                    if key in mapping:
                        values = [cell(row, mapping[key]) for row in rows[index + 1 :] if cell(row, mapping[key])]
                        if values:
                            out.append((target, values[0], "medium", evidence))
                if "样本量" in mapping:
                    numbers = [
                        to_number(cell(row, mapping["样本量"])) for row in rows[index + 1 :]
                    ]
                    numbers = [value for value in numbers if value is not None]
                    if numbers:
                        out.append(("current.sample", numbers[0], "medium", evidence))
                for key in DATE_KEYS:
                    if key in mapping:
                        values = [cell(row, mapping[key]) for row in rows[index + 1 :] if cell(row, mapping[key])]
                        if values:
                            out.append(("current.period", values[0], "medium", evidence))
                            break
                return out
    columns = numeric_columns(rows)
    if len(columns) == 1:
        value_column = columns[0]
        items = []
        for row in rows:
            name = next((item for index, item in enumerate(row) if index != value_column and item), "")
            count = to_number(cell(row, value_column))
            if name and count is not None and not is_total(name) and not name.isdigit():
                items.append({"name": name, "count": int(count) if float(count).is_integer() else count})
        if len(items) >= 3:
            out.append(
                (
                    "current.categories",
                    items,
                    "medium",
                    f"{block.context}（无表头，按“名称+数值”两列推断）",
                )
            )
    return out


def detect_strata(block: Block) -> list[Candidate]:
    rows = block.rows
    header = header_map(rows, STRATA_KEYS + COUNT_KEYS)
    if not header:
        return []
    index, mapping = header
    name_hit = first_key(mapping, STRATA_KEYS)
    value_hit = first_key(mapping, COUNT_KEYS)
    if not name_hit or not value_hit or name_hit[1] == value_hit[1]:
        return []
    items = []
    for row in rows[index + 1 :]:
        name = cell(row, name_hit[1])
        value = to_number(cell(row, value_hit[1]))
        if not name or value is None or is_total(name):
            continue
        items.append({"name": name, "value": value})
    if len(items) < 2:
        return []
    return [("current.strata", items, "high", f"{block.context}（层别表：{name_hit[0]} × {value_hit[0]}）")]


def detect_effect(block: Block) -> list[Candidate]:
    rows = block.rows
    out: list[Candidate] = []
    header = header_map(rows, ("阶段", "时期") + COUNT_KEYS + ("样本量", "总量", "总数") + DATE_KEYS)
    if not header:
        return []
    index, mapping = header
    counts = first_key(mapping, COUNT_KEYS)
    totals = first_key(mapping, ("样本量", "总量", "总数"))
    if not counts:
        return []
    for row in rows[index + 1 :]:
        label = " ".join(row[:3])
        number = to_number(cell(row, counts[1]))
        total = to_number(cell(row, totals[1])) if totals else None
        if number is None:
            continue
        lowered = label.lower()
        is_before = any(word in label for word in ("改善前", "前测", "实施前")) or "before" in lowered
        is_after = any(word in label for word in ("改善后", "后测", "实施后")) or "after" in lowered
        if not (is_before or is_after):
            continue
        prefix = "current" if is_before else "after"
        defects_target = "current.before_defects" if is_before else "after.defects"
        out.append((defects_target, number, "high", f"{block.context}（{label.strip()}）"))
        if total is not None:
            target = "current.sample" if is_before else "after.total"
            out.append((target, total, "high", f"{block.context}（{label.strip()} 样本量）"))
        for key in DATE_KEYS:
            if key in mapping:
                period = cell(row, mapping[key])
                if period:
                    out.append((f"{prefix}.period", period, "medium", f"{block.context}（{label.strip()} 期间）"))
                    break
    return out


def detect_scores(block: Block) -> list[Candidate]:
    rows = block.rows
    header = header_map(rows, SCORE_KEYS)
    if not header:
        return []
    index, mapping = header
    if len(mapping) < 3:
        return []
    ordered = sorted(((key, mapping[key]) for key in SCORE_KEYS if key in mapping), key=lambda item: item[1])
    criteria = [key for key, _ in ordered]
    scores: list[list[float]] = []
    candidates: list[str] = []
    for row in rows[index + 1 :]:
        label = normalise(row[0]) if row else ""
        values = [to_number(cell(row, column)) for _, column in ordered]
        if not label or any(value is None for value in values):
            continue
        candidates.append(label)
        scores.append([float(value) for value in values if value is not None])
    if len(scores) < 2:
        return []
    evidence = f"{block.context}（评分矩阵：{len(criteria)} 维 × {len(scores)} 候选主题）"
    return [
        ("theme.criteria", criteria, "high", evidence),
        ("theme.scores", scores, "high", evidence),
        ("theme.candidates", candidates, "high", evidence),
    ]


def detect_countermeasures(block: Block) -> list[Candidate]:
    rows = block.rows
    header = header_map(
        rows,
        MEASURE_KEYS + CAUSE_KEYS + OWNER_KEYS + WHERE_KEYS + WHEN_KEYS + HOW_KEYS + SCORE_KEYS,
    )
    if not header:
        return []
    index, mapping = header
    measure_hit = first_key(mapping, MEASURE_KEYS)
    if not measure_hit:
        return []
    cause_hit = first_key(mapping, CAUSE_KEYS)
    owner_hit = first_key(mapping, OWNER_KEYS)
    where_hit = first_key(mapping, WHERE_KEYS)
    when_hit = first_key(mapping, WHEN_KEYS)
    how_hit = first_key(mapping, HOW_KEYS)
    score_hits = [(key, mapping[key]) for key in SCORE_KEYS if key in mapping]
    items: list[dict[str, Any]] = []
    for row in rows[index + 1 :]:
        name = cell(row, measure_hit[1])
        if not name or is_total(name):
            continue
        item: dict[str, Any] = {
            "name": name,
            "cause": cell(row, cause_hit[1]) if cause_hit else "",
            "who": cell(row, owner_hit[1]) if owner_hit else "",
            "where": cell(row, where_hit[1]) if where_hit else "",
            "when": cell(row, when_hit[1]) if when_hit else "",
            "how": cell(row, how_hit[1]) if how_hit else "",
        }
        if score_hits:
            scores = {
                key: value
                for key, column in score_hits
                for value in [to_number(cell(row, column))]
                if value is not None
            }
            if scores:
                item["scores"] = scores
        items.append(item)
    if len(items) < 2:
        return []
    five_w = [
        label
        for label, hit in (("真因", cause_hit), ("责任人", owner_hit), ("时间", when_hit), ("做法", how_hit))
        if hit
    ]
    evidence = f"{block.context}（对策表 {len(items)} 条" + (f"；含 {'/'.join(five_w)}" if five_w else "") + "）"
    return [("measures", items, "high", evidence)]


def detect_verification(block: Block) -> list[Candidate]:
    rows = block.rows
    header = header_map(
        rows,
        CAUSE_KEYS
        + VERIFY_KEYS
        + ("维度", "归属", "要因类别", "总分", "评分", "得分", "分数", "影响度", "可控性", "发生频次"),
    )
    if not header:
        return []
    index, mapping = header
    cause_hit = first_key(mapping, CAUSE_KEYS)
    expects = [key for key in VERIFY_KEYS if key in mapping]
    if not cause_hit or not expects:
        return []
    dimension_hit = first_key(mapping, ("维度", "归属", "要因类别"))
    score_hit = first_key(mapping, ("总分", "评分", "得分", "分数", "影响度"))
    items: list[dict[str, Any]] = []
    for row in rows[index + 1 :]:
        name = cell(row, cause_hit[1])
        if not name or is_total(name):
            continue
        item: dict[str, Any] = {"name": name}
        if dimension_hit:
            item["dimension"] = cell(row, dimension_hit[1])
        if score_hit:
            score = to_number(cell(row, score_hit[1]))
            if score is not None:
                item["score"] = score
        for key in expects:
            value = cell(row, mapping[key])
            if key in ("结论", "验证结果"):
                item["result"] = value
            elif key in ("数据来源", "佐证", "证据"):
                item["source"] = value
            elif key == "验证":
                item["method"] = value
        items.append(item)
    if not items:
        return []
    return [("causes", items, "high", f"{block.context}（真因验证表：{cause_hit[0]} + {', '.join(expects)}）")]


def detect_meta(block: Block) -> list[Candidate]:
    text = block.text if block.kind == "text" else "\n".join(" ".join(row) for row in block.rows)
    if not text:
        return []
    out: list[Candidate] = []
    patterns = {
        "meta.topic": ("课题", "主题", "题目"),
        "meta.team": ("圈组", "圈名", "品管圈", "小组"),
        "meta.lead": ("圈长", "组长"),
        "metric": ("指标", "衡量指标"),
        "meta.period": ("活动周期", "活动期间", "改善期间", "活动时间"),
    }
    for target, keys in patterns.items():
        for key in keys:
            match = re.search(rf"{key}\s*[：:]\s*([^\n，,；;。｜|：:]{{2,40}})", text)
            if match:
                value = re.sub(r"[（(][^）)]*[）)]\s*$", "", match.group(1).strip())
                value = trim_meta(value)
                if not value:
                    continue
                out.append((target, value, "medium", f"{block.context}（文本命中「{key}」）"))
                break
    if re.search(r"(越低越好|降低|下降|减少)", text):
        out.append(("meta.direction", "lower", "low", f"{block.context}（改善方向由措辞推断）"))
    return out


def detect_numbers(block: Block) -> list[Candidate]:
    rows = block.rows if block.kind == "table" else [
        [part.strip() for part in line.split("：")] for line in block.text.splitlines() if "：" in line
    ]
    pairs = keyvalue_pairs(rows)
    mapping = {
        "圈能力": "target.capability",
        "样本量": "current.sample",
        "投入人时": "benefit.input_hours",
        "投入费用": "benefit.input_cost",
        "年化节省工时": "benefit.annual_saving_hours",
        "节省工时": "benefit.annual_saving_hours",
        "年化节省费用": "benefit.annual_saving",
        "节省费用": "benefit.annual_saving",
        "回收期": "benefit.payback",
        "计划周数": "plan.planned_weeks",
        "实际周数": "plan.actual_weeks",
        "活动周数": "plan.planned_weeks",
    }
    out: list[Candidate] = []
    for key, value in pairs.items():
        for hint, target in mapping.items():
            if hint not in key:
                continue
            if target == "benefit.payback":
                out.append((target, value, "medium", f"{block.context}（「{key}」）"))
                break
            number = to_number(value)
            if number is not None:
                out.append((target, number, "medium", f"{block.context}（「{key}」）"))
            break
    return out


def detect_intangible(block: Block) -> list[Candidate]:
    rows = block.rows if block.kind == "table" else []
    text = block.text or ""
    haystack = text + " " + " ".join(" ".join(row) for row in rows)
    if not any(word in haystack for word in ("无形成果", "雷达", "圈员能力", "能力成长")):
        return []
    out: list[Candidate] = []
    scale = re.search(r"(\d+\s*[-—~]\s*\d+\s*分制|\d+\s*分制)", haystack)
    if scale:
        out.append(("intangible.scale", re.sub(r"\s+", "", scale.group(1)), "medium", f"{block.context}（量表命中）"))
    dimensions: list[str] = []
    for row in rows:
        label = normalise(row[0]) if row else ""
        if label and len(label) <= 8 and to_number(row[-1] if row else "") is not None:
            if not any(tag in label for tag in ("维度", "项目", "活动前", "活动后", "合计", "评价")):
                dimensions.append(label)
    if len(dimensions) >= 3:
        out.append(("intangible.dimensions", dimensions, "medium", f"{block.context}（维度列）"))
    before = re.search(r"活动前\s*[：:]?\s*(\d+(?:\.\d+)?)", haystack)
    after = re.search(r"活动后\s*[：:]?\s*(\d+(?:\.\d+)?)", haystack)
    if before:
        out.append(("intangible.before_mean", float(before.group(1)), "medium", f"{block.context}（活动前均值）"))
    if after:
        out.append(("intangible.after_mean", float(after.group(1)), "medium", f"{block.context}（活动后均值）"))
    return out


def detect_review(block: Block) -> list[Candidate]:
    text = block.text or "\n".join(" ".join(row) for row in block.rows)
    if not text:
        return []
    out: list[Candidate] = []
    for target, keys in (
        ("review.strengths", ("优点", "做得好的")),
        ("review.weaknesses", ("不足", "待改进")),
        ("review.residual", ("残余问题", "遗留问题", "未解决问题")),
        ("review.next_topic", ("下期主题", "下一期主题", "下阶段主题")),
    ):
        for key in keys:
            match = re.search(rf"{key}\s*[：:]\s*([^\n]{{2,80}})", text)
            if match:
                value: Any = match.group(1).strip()
                if target != "review.next_topic":
                    value = [value]
                out.append((target, value, "medium", f"{block.context}（「{key}」）"))
                break
    return out


def detect_standardization(block: Block) -> list[Candidate]:
    rows = block.rows if block.kind == "table" else []
    if not rows:
        return []
    header = header_map(rows, ("标准化", "作业标准书", "编号", "版本", "生效", "稽核", "训练", "维持"))
    if not header:
        return []
    index, mapping = header
    strong = [
        key
        for key in ("编号", "版本", "生效", "稽核", "训练", "维持", "作业标准书", "标准化")
        if key in mapping
    ]
    header_text = " ".join(rows[index])
    if len(strong) < 2 or any(word in header_text for word in ("要因", "真因", "对策", "总分", "结论")):
        return []
    columns = {key: mapping[key] for key in ("编号", "版本", "生效") if key in mapping}
    documents: list[dict[str, Any]] = []
    for row in rows[index + 1 :]:
        name = normalise(row[0]) if row else ""
        if not name or is_total(name):
            continue
        documents.append(
            {
                "name": name,
                "number": cell(row, columns.get("编号")),
                "version": cell(row, columns.get("版本")),
                "effective": cell(row, columns.get("生效")),
            }
        )
    if not documents:
        return []
    return [("standardization", documents, "medium", f"{block.context}（标准化文件表）")]


DETECTORS: tuple[Callable[[Block], list[Candidate]], ...] = (
    detect_categories,
    detect_strata,
    detect_effect,
    detect_scores,
    detect_countermeasures,
    detect_verification,
    detect_meta,
    detect_numbers,
    detect_intangible,
    detect_review,
    detect_standardization,
)

CONFIDENCE_RANK = {"high": 3, "medium": 2, "low": 1}


# --------------------------------------------------------------------------- #
# minimal draft skeleton
# --------------------------------------------------------------------------- #
def blank_minimal() -> dict:
    return {
        "meta": {"topic": None, "team": None, "lead": None, "period": None, "direction": None},
        "metric": None,
        "current": {
            "period": None,
            "sample": None,
            "before_defects": None,
            "categories": [],
            "strata": [],
            "check_sheet": {"criteria": None, "method": None, "owner": None},
        },
        "target": {"capability": None},
        "after": {"period": None, "defects": None, "total": None},
        "theme": {"candidates": [], "criteria": [], "scores": []},
        "plan": {"planned_weeks": None, "actual_weeks": None, "note": None},
        "causes": [],
        "measures": [],
        "standardization": [],
        "review": {"strengths": [], "weaknesses": [], "residual": [], "next_topic": None},
        "intangible": {"scale": None, "dimensions": [], "before_mean": None, "after_mean": None},
        "benefit": {
            "input_hours": None,
            "input_cost": None,
            "annual_saving_hours": None,
            "annual_saving": None,
            "payback": None,
        },
    }


def get_path(data: dict, path: str) -> Any:
    node: Any = data
    for part in path.split("."):
        if isinstance(node, dict):
            node = node.get(part)
            continue
        if isinstance(node, list) and part.isdigit() and int(part) < len(node):
            node = node[int(part)]
            continue
        return None
    return node


def set_path(data: dict, path: str, value: Any) -> None:
    parts = path.split(".")
    node = data
    for part in parts[:-1]:
        if part not in node or not isinstance(node[part], dict):
            node[part] = {}
        node = node[part]
    node[parts[-1]] = value


REQUIRED_MIN_FIELDS: tuple[tuple[str, str], ...] = (
    ("meta.topic", "课题名称"),
    ("meta.team", "圈组名称"),
    ("meta.lead", "圈长"),
    ("meta.period", "活动周期"),
    ("meta.direction", "指标方向（lower/higher）"),
    ("metric", "指标名称"),
    ("theme.candidates", "候选主题（≥2）"),
    ("theme.criteria", "主题评价维度（≥3）"),
    ("theme.scores", "主题评分矩阵"),
    ("plan.planned_weeks", "计划周数"),
    ("plan.actual_weeks", "实际周数"),
    ("current.period", "现状收集期间"),
    ("current.sample", "现状样本量"),
    ("current.categories", "问题类别与频次（≥3，降序）"),
    ("current.strata", "层别数据"),
    ("current.before_defects", "改善前异常例数"),
    ("target.capability", "圈能力（%）"),
    ("after.period", "改善后收集期间"),
    ("after.defects", "改善后异常例数"),
    ("causes", "已验证真因（含维度/评分/来源/结果）"),
    ("measures", "对策（含真因映射 + 5W1H + 评分）"),
    ("standardization", "标准化文件（编号/版本/生效/稽核）"),
    ("review.strengths", "优点"),
    ("review.weaknesses", "不足"),
    ("review.residual", "残余问题"),
    ("review.next_topic", "下期主题"),
    ("intangible.scale", "无形成果量表"),
    ("intangible.dimensions", "无形成果维度"),
    ("intangible.before_mean", "活动前均值"),
    ("intangible.after_mean", "活动后均值"),
    ("benefit.input_hours", "投入人时"),
    ("benefit.input_cost", "投入费用"),
    ("benefit.annual_saving", "年化节省费用"),
    ("benefit.payback", "回收期"),
)


def empty(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return not value.strip()
    if isinstance(value, (list, dict)):
        return len(value) == 0
    return False


def file_hash(path: Path, limit: int = 2_000_000) -> str:
    digest = hashlib.sha1()
    read = 0
    with path.open("rb") as handle:
        while read < limit:
            chunk = handle.read(65536)
            if not chunk:
                break
            digest.update(chunk)
            read += len(chunk)
    return digest.hexdigest()[:12]


def iter_files(root: Path) -> list[Path]:
    files: list[Path] = []
    for path in sorted(root.rglob("*")):
        if path.is_dir():
            continue
        if any(part in SKIP_DIRS or part.startswith("~$") for part in path.parts):
            continue
        if path.name.startswith("."):
            continue
        files.append(path)
    return files


def detect_block(block: Block) -> list[Candidate]:
    found: list[Candidate] = []
    for detector in DETECTORS:
        try:
            found.extend(detector(block))
        except Exception as error:  # noqa: BLE001 - one bad detector must not kill the scan
            found.append(("__detector_error__", f"{detector.__name__}: {error}", "low", block.context))
    return found


def scan(root: Path) -> dict:
    files = iter_files(root)
    inventory: list[dict[str, Any]] = []
    blocks: list[Block] = []
    for path in files:
        rel = path.relative_to(root).as_posix()
        try:
            stat = path.stat()
        except OSError:
            continue
        inventory.append(
            {
                "path": rel,
                "suffix": path.suffix.lower(),
                "size": stat.st_size,
                "sha1_12": file_hash(path) if stat.st_size else "",
                "readable": path.suffix.lower() in SUPPORTED,
            }
        )
        blocks.extend(extract_file(path, rel))

    candidates: list[dict[str, Any]] = []
    for block in blocks:
        for target, value, confidence, evidence in detect_block(block):
            candidates.append(
                {"target": target, "value": value, "confidence": confidence, "evidence": evidence}
            )

    best: dict[str, dict[str, Any]] = {}
    for candidate in candidates:
        target = candidate["target"]
        if target.startswith("__"):
            continue
        current = best.get(target)
        if current is None or CONFIDENCE_RANK.get(candidate["confidence"], 0) > CONFIDENCE_RANK.get(
            current["confidence"], 0
        ):
            best[target] = candidate

    draft = blank_minimal()
    for target, candidate in best.items():
        set_path(draft, target, candidate["value"])

    missing = [
        {"path": path, "label": label}
        for path, label in REQUIRED_MIN_FIELDS
        if empty(get_path(draft, path))
    ]
    return {
        "root": str(root),
        "inventory": inventory,
        "blocks": [block.to_json() for block in blocks],
        "candidates": candidates,
        "resolved": best,
        "draft": draft,
        "missing": missing,
    }


# --------------------------------------------------------------------------- #
# reporting
# --------------------------------------------------------------------------- #
EXTRACTION_PROMPT = """你是品管圈（QCC）数据抽取助手。只做“读证据、填字段”，不要编造。

输入：
1. 证据索引：{evidence}
2. 待填草稿：{draft}
3. 必备字段说明：docs/qcc_data_requirements.md

任务：
- 逐条阅读证据索引里的表格（rows）与文本（text），把每一项必备数据补进草稿 YAML；
- 每个填入的值都要注明来源（文件名 + 工作表/页/幻灯片 + 行号或单元格）；
- 找不到的字段保持 null，并列为“缺失，需要使用者补充”；不得用估计值或示例值填充；
- 同一指标出现两个不同数值时，两处都列出请人工裁决，不要自行取平均；
- “合计/小计/平均”行不要当作事实数据。

输出：
- 更新后的 qcc-min-data.yaml
- 抽取说明 reports/data-extraction-notes.md：字段 → 值 → 来源 → 置信度（高/中/低）
- 缺失清单：字段 → 为什么需要 → 建议向使用者索要什么
"""


def format_report(result: dict, draft_path: Path, evidence_path: Path) -> str:
    inventory = result["inventory"]
    by_suffix: dict[str, int] = {}
    for item in inventory:
        key = item["suffix"] or "(无扩展名)"
        by_suffix[key] = by_suffix.get(key, 0) + 1

    lines = [
        "# QCC 数据目录扫描报告",
        "",
        f"- 扫描目录：`{result['root']}`",
        f"- 文件数：{len(inventory)}（{', '.join(f'{k}×{v}' for k, v in sorted(by_suffix.items()))}）",
        f"- 已识别字段：{len(result['resolved'])} 项；仍缺失：{len(result['missing'])} 项",
        "",
        "## 1. 文件清单",
        "",
        "| 文件 | 类型 | 大小 | 可解析 |",
        "|---|---|---:|---|",
    ]
    for item in inventory:
        lines.append(
            f"| `{item['path']}` | {item['suffix'] or '-'} | {item['size']} B | "
            f"{'是' if item['readable'] else '否'} |"
        )

    notes = [block for block in result["blocks"] if block.get("note")]
    errors = [candidate for candidate in result["candidates"] if candidate["target"].startswith("__")]
    if notes or errors:
        lines += ["", "## 2. 需要处理的情况", ""]
        for block in notes:
            lines.append(f"- `{block['source']}`（{block['locator']}）：{block['note']}")
        for item in errors:
            lines.append(f"- 检测器异常 `{item['evidence']}`：{item['value']}（需修复脚本）")

    lines += ["", "## 3. 已识别到的数据", "", "| 目标字段 | 值（摘要） | 置信度 | 证据 |", "|---|---|---|---|"]
    if result["resolved"]:
        for target, candidate in sorted(result["resolved"].items()):
            summary = json.dumps(candidate["value"], ensure_ascii=False).replace("|", "\\|")
            if len(summary) > 90:
                summary = summary[:87] + "..."
            lines.append(
                f"| `{target}` | {summary} | {candidate['confidence']} | {candidate['evidence']} |"
            )
    else:
        lines.append("| - | 目录里没有找到可识别的表格/文本 | - | - |")

    lines += ["", "## 4. 仍缺失（需使用者补充或模型进一步抽取）", ""]
    lines += [f"- `{item['path']}` — {item['label']}" for item in result["missing"]] or ["- 无"]

    lines += [
        "",
        "## 5. 下一步",
        "",
        "```bash",
        "# 1) 让模型按证据填充草稿（或人工补）",
        f"#    证据索引：{evidence_path}",
        f"#    待填草稿：{draft_path}",
        "#    字段说明：docs/qcc_data_requirements.md",
        "# 2) 推导 + 校验 + 统计复算（一条命令）",
        f"python scripts/qcc_pipeline.py --min {draft_path} --workspace qcc-workspace",
        "# 3) 出片后合并合规校验",
        "python scripts/check_qcc_method_compliance.py qcc-workspace/output/qcc-review-ready.pptx "
        "--data qcc-workspace/input/qcc-data.yaml",
        "```",
        "",
        "## 6. 交给模型的抽取指令",
        "",
        "```text",
        EXTRACTION_PROMPT.format(evidence=evidence_path, draft=draft_path).strip(),
        "```",
        "",
    ]
    return "\n".join(lines)


def main() -> int:
    parser = argparse.ArgumentParser(description="Scan a folder of user files into QCC intake evidence.")
    parser.add_argument("--dir", dest="root", type=Path, required=True, help="使用者的原始数据文件夹")
    parser.add_argument("--workspace", type=Path, default=Path("qcc-workspace"))
    parser.add_argument("--draft", type=Path, default=None)
    parser.add_argument("--report", type=Path, default=None)
    parser.add_argument("--evidence", type=Path, default=None)
    parser.add_argument("--json", action="store_true", help="把识别结果打印到标准输出")
    args = parser.parse_args()

    if not args.root.exists() or not args.root.is_dir():
        print(f"目录不存在或不是文件夹：{args.root}", file=sys.stderr)
        return 2
    if yaml is None:
        print("需要 PyYAML：pip install PyYAML", file=sys.stderr)
        return 2

    result = scan(args.root)
    if not result["inventory"]:
        print(f"目录里没有文件：{args.root}", file=sys.stderr)
        return 1

    draft_path = args.draft or (args.workspace / "input" / "qcc-min-data.draft.yaml")
    report_path = args.report or (args.workspace / "reports" / "data-scan-report.md")
    evidence_path = args.evidence or (args.workspace / "reports" / "data-scan-evidence.json")
    for path in (draft_path, report_path, evidence_path):
        path.parent.mkdir(parents=True, exist_ok=True)

    header = (
        "# 由 qcc_scan_inputs.py 生成的草稿：只包含在证据中确实找到的值。\n"
        "# 其余字段为 null，需按 reports/data-scan-report.md 补全，不得编造。\n"
    )
    draft_path.write_text(
        header + yaml.safe_dump(result["draft"], allow_unicode=True, sort_keys=False),
        encoding="utf-8",
    )
    evidence_path.write_text(
        json.dumps(
            {
                "root": result["root"],
                "inventory": result["inventory"],
                "blocks": result["blocks"],
                "candidates": result["candidates"],
                "resolved": {
                    target: {
                        "value": item["value"],
                        "confidence": item["confidence"],
                        "evidence": item["evidence"],
                    }
                    for target, item in result["resolved"].items()
                },
                "missing": result["missing"],
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    report = format_report(result, draft_path, evidence_path)
    report_path.write_text(report, encoding="utf-8")

    print(report)
    print(f"\n草案：{draft_path}")
    print(f"报告：{report_path}")
    print(f"证据索引：{evidence_path}")

    if args.json:
        print(json.dumps({"resolved": result["resolved"], "missing": result["missing"]}, ensure_ascii=False, indent=2))

    return 0 if result["resolved"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
